import os
import re
import time
import uuid
import tempfile
import subprocess
import sys
import shutil
import hashlib
from pathlib import Path
from urllib.parse import urlparse, urljoin, urldefrag

import streamlit as st
from faster_whisper import WhisperModel
from dotenv import load_dotenv
from google import genai
from google.genai import types

from pypdf import PdfReader
from docx import Document

# SQL persistence
from database import (
    create_session,
    load_messages,
    save_message,
    save_source,
    clear_session_messages,
    touch_session,
    list_sessions,
)

import pymupdf  # PyMuPDF
import requests
from bs4 import BeautifulSoup


# ============================================================
# OPTIONAL PLAYWRIGHT
# ============================================================

try:
    from playwright.sync_api import sync_playwright

    PLAYWRIGHT_AVAILABLE = True

except Exception:

    PLAYWRIGHT_AVAILABLE = False


# ============================================================
# ENVIRONMENT
# SERVER-SAFE / UTF-8 SAFE
# ============================================================

def load_environment_safely():
    """
    Load .env without depending on the Windows server's default
    cp1252/charmap encoding.

    Primary: UTF-8
    Fallback: Latin-1 for legacy Windows-encoded .env files.
    """
    try:
        load_dotenv(encoding="utf-8", override=False)
        return
    except UnicodeDecodeError:
        pass

    try:
        load_dotenv(encoding="latin-1", override=False)
        return
    except Exception as error:
        raise RuntimeError(
            "Unable to read the .env file. Please save .env as UTF-8 and try again."
        ) from error


load_environment_safely()

# Prefer UTF-8 for Python standard text streams when supported.
try:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

GEMINI_API_KEY = os.getenv(
    "GEMINI_API_KEY",
    ""
).strip()

WHISPER_MODEL_NAME = os.getenv(
    "WHISPER_MODEL",
    "medium"
).strip()

GEMINI_MODEL = os.getenv(
    "GEMINI_MODEL",
    "gemini-3.5-flash"
).strip()

if not GEMINI_MODEL:
    GEMINI_MODEL = "gemini-3.5-flash"

FALLBACK_MODELS = [
    "gemini-3.5-flash-lite",
]


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Multilingual AI Voice Assistant",
    page_icon="🎙️",
    layout="wide",
)


# ============================================================
# API KEY
# ============================================================

if not GEMINI_API_KEY:

    st.error(
        "❌ GEMINI_API_KEY is missing.\n\n"
        "Please add GEMINI_API_KEY to your .env file."
    )

    st.stop()


# ============================================================
# GEMINI CLIENT
# ============================================================

client = genai.Client(
    api_key=GEMINI_API_KEY
)


# ============================================================
# LANGUAGES
# ============================================================

LANGUAGE_NAMES = {

    "en": "English",
    "hi": "Hindi",
    "bn": "Bengali",
    "or": "Odia",
    "ta": "Tamil",
    "te": "Telugu",
    "ne": "Nepali",
    "gu": "Gujarati",
    "as": "Assamese",
    "mr": "Marathi",
    "kn": "Kannada",
    "ml": "Malayalam",
    "pa": "Punjabi",
    "ur": "Urdu",

    "fr": "French",
    "de": "German",
    "es": "Spanish",
    "it": "Italian",
    "pt": "Portuguese",

    "ja": "Japanese",
    "ko": "Korean",
    "zh": "Chinese",

    "ru": "Russian",
    "ar": "Arabic",
    "tr": "Turkish",

    "id": "Indonesian",
    "vi": "Vietnamese",
    "th": "Thai",

    "pl": "Polish",
    "nl": "Dutch",
    "sv": "Swedish",
    "da": "Danish",
    "fi": "Finnish",

    "uk": "Ukrainian",
    "cs": "Czech",
    "ro": "Romanian",
    "el": "Greek",
    "he": "Hebrew",

    "hu": "Hungarian",
    "no": "Norwegian",
    "sk": "Slovak",
    "bg": "Bulgarian",
    "hr": "Croatian",
    "sr": "Serbian",
    "sl": "Slovenian",

    "sw": "Swahili",
    "af": "Afrikaans",
}


# ============================================================
# EDGE TTS VOICES
# ============================================================

TTS_VOICES = {

    "en": "en-US-AriaNeural",

    "hi": "hi-IN-SwaraNeural",
    "bn": "bn-IN-TanishaaNeural",
    "or": "or-IN-SubhasiniNeural",
    "ta": "ta-IN-PallaviNeural",
    "te": "te-IN-ShrutiNeural",
    "ne": "ne-NP-HemkalaNeural",
    "gu": "gu-IN-DhwaniNeural",
    "as": "as-IN-PriyomNeural",
    "mr": "mr-IN-AarohiNeural",
    "kn": "kn-IN-SapnaNeural",
    "ml": "ml-IN-SobhanaNeural",
    "pa": "pa-IN-OjasNeural",
    "ur": "ur-PK-AsadNeural",

    "fr": "fr-FR-DeniseNeural",
    "de": "de-DE-KatjaNeural",
    "es": "es-ES-ElviraNeural",
    "it": "it-IT-ElsaNeural",
    "pt": "pt-BR-FranciscaNeural",

    "ja": "ja-JP-NanamiNeural",
    "ko": "ko-KR-SunHiNeural",
    "zh": "zh-CN-XiaoxiaoNeural",

    "ru": "ru-RU-SvetlanaNeural",
    "ar": "ar-SA-ZariyahNeural",
    "tr": "tr-TR-EmelNeural",

    "id": "id-ID-GadisNeural",
    "vi": "vi-VN-HoaiMyNeural",
    "th": "th-TH-PremwadeeNeural",

    "pl": "pl-PL-ZofiaNeural",
    "nl": "nl-NL-ColetteNeural",
    "sv": "sv-SE-SofieNeural",
    "da": "da-DK-ChristelNeural",
    "fi": "fi-FI-NooraNeural",

    "uk": "uk-UA-PolinaNeural",
    "cs": "cs-CZ-VlastaNeural",
    "ro": "ro-RO-AlinaNeural",
    "el": "el-GR-AthinaNeural",
    "he": "he-IL-HilaNeural",

    "hu": "hu-HU-NoemiNeural",
    "no": "nb-NO-IselinNeural",
    "sk": "sk-SK-ViktoriaNeural",
    "bg": "bg-BG-KalinaNeural",
    "hr": "hr-HR-GabrijelaNeural",
    "sr": "sr-RS-SophieNeural",
    "sl": "sl-SI-PetraNeural",

    "sw": "sw-KE-ZuriNeural",
    "af": "af-ZA-AdriNeural",
}


# ============================================================
# ROMANIZED BENGALI
# ============================================================

ROMAN_BENGALI = {

    "ami", "amar", "amake", "amra", "amader",
    "tumi", "tomar", "tomake", "tomra", "tomader",
    "apni", "apnar", "apnake",

    "se", "she", "tar", "take",

    "ki", "ke", "kake", "kemon", "kamon",
    "keno", "kothay", "kotha", "kokhon",
    "kivabe", "kirokom",

    "achi", "achhi", "acho", "achho", "ache", "achen",
    "chilam", "chhilam", "chilo", "chhilo",

    "korchi", "korchhi", "korcho", "korchho",
    "korchen", "korbo", "korbe", "koro",
    "korben", "kore",

    "jacchi", "jachhi", "jacchhi", "jachcho",
    "jaccho", "jabo", "jabe",

    "gechi", "gechhi", "geche",

    "asche", "aschhe", "aschhi", "aschi",
    "ashchi", "ashche",

    "khacchi", "khachhi", "khaccho",
    "khachcho", "khabo", "kheye",

    "bhalo", "valo", "valobasha", "bhalobasha",
    "kharap", "sundor", "onek", "khub", "ektu",
    "sob", "shob", "kichu", "kono",

    "ekhane", "okhane", "sekhane", "sekhaney",

    "aj", "aaj", "kal", "ekhon",

    "naam", "nam",

    "hobe", "hoy", "hoye", "hoyeche",

    "dorkar", "proyojon",

    "chai", "chao", "chaichi",

    "dao", "den", "de",

    "dekh", "dekho",

    "bolo", "bol", "bolchi", "bolcho",

    "jante", "jani", "janina",

    "parbo", "pari", "parena",

    "na", "nei", "noy",

    "hya", "ha", "haan",

    "dhonnobad",

    "bari", "ghor", "bondhu",
    "ma", "baba", "dada", "didi",
    "bhai", "bon",

    "porashona", "porchi", "porte",

    "kaj", "chakri", "project", "office",
}


# ============================================================
# ROMANIZED HINDI
# ============================================================

ROMAN_HINDI = {

    "main", "mera", "meri", "mujhe",
    "hum", "hamara", "hamari",

    "aap", "aapka", "aapki",
    "tum", "tumhara", "tumhe",

    "kya", "kaise", "kaisa", "kyun",
    "kyon", "kahan", "kab", "kaun",

    "hai", "hain", "hoon",

    "tha", "thi", "the",

    "kar", "karo", "karna",
    "karunga", "karenge",

    "ja", "jana", "jaunga", "jaoge", "jao",

    "aa", "aana", "aunga",

    "acha", "accha", "achha",

    "bahut", "thoda", "sab", "kuch",

    "ghar", "dost",

    "chahiye",

    "nahi", "nahin",

    "dhanyavad", "shukriya",
}


# ============================================================
# OTHER LATIN LANGUAGE WORDS
# ============================================================

ROMAN_LANGUAGE_WORDS = {

    "fr": {
        "bonjour", "salut", "merci", "comment",
        "vous", "êtes", "suis", "avec", "pour",
        "dans", "une", "des", "les", "est",
        "oui", "non", "je", "tu", "il", "elle",
        "nous", "mon", "ma", "mes", "très"
    },

    "de": {
        "hallo", "danke", "bitte", "wie", "geht",
        "dir", "ihnen", "ich", "du", "er", "sie",
        "wir", "mein", "meine", "nicht", "ja",
        "nein", "und", "ist", "das", "die", "der",
        "mit", "für", "von", "auf", "sehr"
    },

    "es": {
        "hola", "gracias", "como", "cómo", "estas",
        "estás", "usted", "tú", "yo", "nosotros",
        "ellos", "ella", "que", "qué", "para",
        "por", "con", "una", "uno", "los", "las",
        "del", "es", "muy", "bien", "sí", "si", "no"
    },

    "it": {
        "ciao", "grazie", "come", "stai", "sta",
        "sono", "sei", "io", "tu", "lui", "lei",
        "noi", "che", "chi", "cosa", "per",
        "con", "una", "uno", "gli", "le",
        "non", "si", "sì", "bene", "molto"
    },

    "pt": {
        "olá", "ola", "obrigado", "obrigada", "como",
        "está", "voce", "você", "eu", "tu",
        "ele", "ela", "nos", "nós", "que",
        "para", "por", "com", "uma", "um",
        "não", "sim", "bem", "muito"
    },

    "tr": {
        "merhaba", "selam", "teşekkür", "tesekkur",
        "nasılsın", "nasilsin", "nasıl", "nasil",
        "ben", "sen", "siz", "biz", "bu",
        "ne", "neden", "nerede", "evet", "hayır",
        "hayir", "değil", "degil", "çok", "cok"
    },

    "id": {
        "halo", "hai", "terima", "kasih", "bagaimana",
        "kamu", "anda", "saya", "aku", "dia",
        "kami", "kita", "apa", "kenapa", "dimana",
        "di", "yang", "dan", "untuk", "dengan",
        "tidak", "iya", "ya", "baik", "sangat"
    },

    "vi": {
        "xin", "chào", "chao", "cảm", "cam",
        "ơn", "on", "không", "khong", "bạn",
        "ban", "tôi", "toi", "mình", "minh",
        "là", "la", "gì", "gi", "nào", "nao",
        "ở", "o", "đâu", "dau", "vâng", "vang"
    },

    "nl": {
        "hallo", "hoi", "bedankt", "dank", "hoe",
        "gaat", "het", "met", "jou", "u",
        "ik", "jij", "je", "hij", "zij",
        "wij", "wat", "waar", "waarom", "niet",
        "ja", "nee", "goed", "voor", "van"
    },

    "sv": {
        "hej", "tack", "hur", "mår", "mar",
        "du", "jag", "han", "hon", "vi",
        "vad", "var", "varför", "varfor",
        "inte", "ja", "nej", "bra", "och",
        "det", "är", "ar"
    },

    "da": {
        "hej", "tak", "hvordan", "har", "du",
        "det", "jeg", "han", "hun", "vi",
        "hvad", "hvor", "hvorfor", "ikke",
        "ja", "nej", "godt", "og", "er"
    },

    "fi": {
        "hei", "kiitos", "mitä", "mita", "kuuluu",
        "sinä", "sina", "minä", "mina", "hän",
        "han", "me", "te", "miksi", "missä",
        "missa", "ei", "kyllä", "kylla", "hyvä",
        "hyva"
    },

    "pl": {
        "cześć", "czesc", "dzień", "dzien", "dobry",
        "dziękuję", "dziekuje", "jak", "się", "sie",
        "masz", "mam", "jest", "nie", "tak",
        "co", "gdzie", "dlaczego", "ja", "ty",
        "my", "wy", "oni"
    },

    "ro": {
        "salut", "bună", "buna", "mulțumesc",
        "multumesc", "cum", "ești", "esti",
        "sunt", "eu", "tu", "el", "ea",
        "noi", "ce", "unde", "de", "nu",
        "da", "bine", "foarte", "pentru", "cu"
    },

    "hu": {
        "szia", "köszönöm", "koszonom", "hogy",
        "vagy", "én", "en", "te", "ő", "o",
        "mi", "ti", "ők", "ok", "nem", "igen",
        "miért", "miert", "hol", "jó", "jo"
    },

    "cs": {
        "ahoj", "děkuji", "dekuji", "jak", "se",
        "máš", "mas", "mám", "mam", "jsem",
        "jsi", "on", "ona", "my", "co",
        "kde", "proč", "proc", "ano", "ne",
        "dobře", "dobre"
    },

    "sk": {
        "ahoj", "ďakujem", "dakujem", "ako",
        "sa", "máš", "mas", "som", "si",
        "on", "ona", "my", "čo", "co",
        "kde", "prečo", "preco", "áno", "ano",
        "nie", "dobre"
    },

    "sl": {
        "zdravo", "hvala", "kako", "si", "kaj",
        "jaz", "ti", "on", "ona", "mi",
        "kje", "zakaj", "ne", "da", "dobro"
    },

    "hr": {
        "bok", "zdravo", "hvala", "kako", "si",
        "ja", "ti", "on", "ona", "mi",
        "što", "sto", "gdje", "zasto", "zašto",
        "da", "ne", "dobro"
    },

    "sr": {
        "zdravo", "ćao", "cao", "hvala", "kako",
        "si", "ja", "ti", "on", "ona", "mi",
        "šta", "sta", "gde", "zasto", "zašto",
        "da", "ne", "dobro"
    },

    "sw": {
        "habari", "asante", "tafadhali", "jina",
        "langu", "wewe", "mimi", "yeye", "sisi",
        "nini", "wapi", "kwa", "na", "ni",
        "hapana", "ndiyo", "nzuri", "sana"
    },

    "af": {
        "hallo", "dankie", "hoe", "gaan", "dit",
        "met", "jou", "ek", "jy", "hy",
        "sy", "ons", "wat", "waar", "hoekom",
        "nie", "ja", "nee", "goed", "baie"
    },

    "no": {
        "hei", "hallo", "takk", "hvordan", "har",
        "du", "jeg", "han", "hun", "vi",
        "hva", "hvor", "hvorfor", "ikke",
        "ja", "nei", "bra", "og", "er"
    },
}


# ============================================================
# SCRIPT DETECTION
# ============================================================

def detect_script_language(text):

    if not text:
        return None

    counts = {}

    for ch in text:

        if not ch.isalpha():
            continue

        code = ord(ch)
        language = None

        if 0x0980 <= code <= 0x09FF:
            language = "bn"

        elif 0x0900 <= code <= 0x097F:
            language = "hi"

        elif 0x0A00 <= code <= 0x0A7F:
            language = "pa"

        elif 0x0A80 <= code <= 0x0AFF:
            language = "gu"

        elif 0x0B00 <= code <= 0x0B7F:
            language = "or"

        elif 0x0B80 <= code <= 0x0BFF:
            language = "ta"

        elif 0x0C00 <= code <= 0x0C7F:
            language = "te"

        elif 0x0C80 <= code <= 0x0CFF:
            language = "kn"

        elif 0x0D00 <= code <= 0x0D7F:
            language = "ml"

        elif 0x0E00 <= code <= 0x0E7F:
            language = "th"

        elif 0x0590 <= code <= 0x05FF:
            language = "he"

        elif (
            0x0600 <= code <= 0x06FF
            or 0x0750 <= code <= 0x077F
            or 0x08A0 <= code <= 0x08FF
            or 0xFB50 <= code <= 0xFDFF
            or 0xFE70 <= code <= 0xFEFF
        ):
            language = "ar"

        elif (
            0x3040 <= code <= 0x309F
            or 0x30A0 <= code <= 0x30FF
            or 0x31F0 <= code <= 0x31FF
        ):
            language = "ja"

        elif 0xAC00 <= code <= 0xD7AF:
            language = "ko"

        elif 0x4E00 <= code <= 0x9FFF:
            language = "zh"

        elif 0x0400 <= code <= 0x04FF:
            language = "ru"

        elif 0x0370 <= code <= 0x03FF:
            language = "el"

        if language:

            counts[language] = (
                counts.get(language, 0) + 1
            )

    if not counts:
        return None

    strongest = max(
        counts,
        key=counts.get
    )

    if counts[strongest] >= 2:
        return strongest

    return None


# ============================================================
# URDU DETECTION
# ============================================================

def contains_urdu_specific_characters(text):

    urdu_specific = set(
        "ٹڈڑںھہےےژچگپ"
    )

    return any(
        ch in urdu_specific
        for ch in text
    )


# ============================================================
# NORMALIZE
# ============================================================

def normalize_words(text):

    text = str(text).lower().strip()

    replacements = {
        "’": "'",
        "‘": "'",
        "“": '"',
        "”": '"',
    }

    for old, new in replacements.items():
        text = text.replace(old, new)

    return re.findall(
        r"[^\W_]+",
        text,
        flags=re.UNICODE
    )


# ============================================================
# ROMAN LANGUAGE DETECTION
# ============================================================

def detect_roman_language(text):

    if not text:
        return "en"

    words = normalize_words(text)

    if not words:
        return "en"

    word_set = set(words)

    normalized = " ".join(words)

    bengali_phrases = [

        "tumi kemon acho",
        "tumi kamon acho",
        "tumi kemon achho",
        "tumi kamon achho",

        "ami bhalo achi",
        "ami valo achi",
        "ami bhalo achhi",
        "ami valo achhi",

        "tumi ki korcho",
        "tumi ki korchho",
        "tumi ki korbe",

        "kothay jabe",
        "kivabe korbo",

        "amar naam",
        "amar nam",

        "ami jani na",

        "tumi ki",
        "amar ki",
        "eta ki",
        "ota ki",

        "ki korcho",
        "ki korchho",

        "kemon acho",
        "kamon acho",
    ]

    for phrase in bengali_phrases:

        if phrase in normalized:
            return "bn"

    bn_score = len(
        word_set.intersection(
            ROMAN_BENGALI
        )
    )

    strong_bn = {

        "ami", "amar", "amake",
        "tumi", "tomar", "tomake",
        "apni", "apnar",

        "kemon", "kamon",
        "kothay", "kivabe",

        "achi", "achhi",
        "acho", "achho",

        "korchi", "korchhi",
        "korcho", "korchho",

        "korbo", "korbe",

        "jacchi", "jachhi",
        "jabo", "jabe",

        "bhalo", "valo",

        "ekhane", "okhane",

        "bolchi", "bolcho",

        "jani", "janina",

        "dhonnobad",
    }

    strong_bn_score = len(
        word_set.intersection(
            strong_bn
        )
    )

    if strong_bn_score >= 1:
        return "bn"

    hindi_phrases = [

        "main theek hoon",
        "main thik hoon",

        "aap kaise hain",
        "tum kaise ho",

        "tum kya kar rahe ho",

        "mujhe nahi pata",
        "mujhe kya karna hai",

        "kahan ja rahe ho",

        "kaise ho",
    ]

    for phrase in hindi_phrases:

        if phrase in normalized:
            return "hi"

    hi_score = len(
        word_set.intersection(
            ROMAN_HINDI
        )
    )

    strong_hi = {

        "main", "mera", "meri", "mujhe",

        "hum", "hamara", "hamari",

        "aap", "aapka", "aapki",

        "tum", "tumhara", "tumhe",

        "kya", "kaise", "kaisa",

        "kyun", "kyon", "kahan",

        "hai", "hain", "hoon",

        "acha", "accha", "achha",

        "bahut", "nahi", "nahin",

        "chahiye",

        "dhanyavad", "shukriya",
    }

    strong_hi_score = len(
        word_set.intersection(
            strong_hi
        )
    )

    if (
        strong_hi_score >= 2
        and strong_hi_score >= bn_score
    ):
        return "hi"

    scores = {}

    for language_code, vocabulary in (
        ROMAN_LANGUAGE_WORDS.items()
    ):

        score = len(
            word_set.intersection(
                vocabulary
            )
        )

        if score:
            scores[language_code] = score

    distinctive = {

        "bonjour": "fr",
        "merci": "fr",
        "salut": "fr",

        "danke": "de",
        "bitte": "de",

        "hola": "es",
        "gracias": "es",

        "ciao": "it",
        "grazie": "it",

        "olá": "pt",
        "ola": "pt",
        "obrigado": "pt",
        "obrigada": "pt",

        "merhaba": "tr",
        "nasılsın": "tr",
        "nasilsin": "tr",

        "habari": "sw",
        "asante": "sw",

        "szia": "hu",
        "ahoj": "cs",

        "hei": "no",
        "takk": "no",
    }

    for word in words:

        if word in distinctive:
            return distinctive[word]

    if scores:

        ranked = sorted(
            scores.items(),
            key=lambda x: x[1],
            reverse=True
        )

        best_language, best_score = ranked[0]

        second_score = (
            ranked[1][1]
            if len(ranked) > 1
            else 0
        )

        if best_score >= 2:

            if (
                best_score > second_score
                or best_score >= 3
            ):
                return best_language

    english_words = {

        "the", "is", "are", "am",
        "you", "your", "how",
        "what", "why", "where",
        "when", "who",
        "can", "could",
        "would", "should",
        "please", "help",
        "hello", "hi", "hey",
        "thanks", "thank",
        "good", "morning",
        "evening", "today",
        "tomorrow", "yesterday",
        "this", "that", "with",
        "from", "for", "and",
        "but", "not", "have",
        "has", "had",
        "do", "does", "did",
        "want", "need",
        "know", "tell",
        "give", "make",
        "computer", "project",
        "code", "error",
        "python",
    }

    english_score = len(
        word_set.intersection(
            english_words
        )
    )

    if english_score >= 1:
        return "en"

    return "en"


# ============================================================
# TEXT LANGUAGE DETECTION
# ============================================================

def detect_text_language(text):

    if not text:
        return "en"

    text = str(text).strip()

    if not text:
        return "en"

    script_language = detect_script_language(
        text
    )

    if script_language:

        if script_language == "ar":

            if contains_urdu_specific_characters(
                text
            ):
                return "ur"

            return "ar"

        return script_language

    return detect_roman_language(
        text
    )


# ============================================================
# WHISPER
# ============================================================

@st.cache_resource(show_spinner=False)
def load_whisper():

    return WhisperModel(
        WHISPER_MODEL_NAME,
        device="cpu",
        compute_type="int8",
        cpu_threads=4,
        num_workers=1,
    )


# ============================================================
# FFMPEG CHECK
# ============================================================

def check_ffmpeg():

    return shutil.which(
        "ffmpeg"
    ) is not None


# ============================================================
# AUDIO CONVERSION
# ============================================================

def convert_audio_to_wav(
    input_path
):

    if not check_ffmpeg():

        raise RuntimeError(
            "FFmpeg was not found. "
            "Please install FFmpeg and add it to PATH."
        )

    output_path = os.path.join(

        tempfile.gettempdir(),

        f"converted_{uuid.uuid4().hex}.wav"
    )

    command = [

        "ffmpeg",
        "-y",

        "-i",
        input_path,

        "-ar",
        "16000",

        "-ac",
        "1",

        "-c:a",
        "pcm_s16le",

        output_path,
    ]

    try:

        result = subprocess.run(

            command,

            stdout=subprocess.PIPE,

            stderr=subprocess.PIPE,

            text=True,

            timeout=120,
        )

        if result.returncode != 0:

            raise RuntimeError(
                result.stderr[-3000:]
            )

        if not os.path.exists(
            output_path
        ):

            raise RuntimeError(
                "FFmpeg did not create the WAV file."
            )

        return output_path

    except FileNotFoundError:

        raise RuntimeError(
            "FFmpeg was not found."
        )


# ============================================================
# TRANSCRIBE AUDIO
# ============================================================

def transcribe_audio(
    model,
    audio_file,
    return_segments=False,
):

    try:

        segments, info = model.transcribe(

            audio_file,

            language=None,

            task="transcribe",

            beam_size=5,

            best_of=5,

            temperature=0,

            condition_on_previous_text=False,

            vad_filter=True,

            vad_parameters={
                "min_silence_duration_ms": 500,
            },
        )

        segments = list(
            segments
        )

        segment_data = []

        for segment in segments:

            text = segment.text.strip()

            if text:

                segment_data.append({

                    "start": float(
                        segment.start
                    ),

                    "end": float(
                        segment.end
                    ),

                    "text": text,
                })

        text = " ".join(

            item["text"]

            for item in segment_data

        ).strip()

        if not text:

            return (
                "en",
                "",
                []
            ) if return_segments else (
                "en",
                ""
            )

        whisper_language = str(
            getattr(
                info,
                "language",
                "en"
            )
        ).strip().lower()

        whisper_probability = float(
            getattr(
                info,
                "language_probability",
                0.0
            ) or 0.0
        )

        whisper_map = {
            "bh": "hi",
        }

        whisper_language = whisper_map.get(
            whisper_language,
            whisper_language
        )

        if whisper_language not in LANGUAGE_NAMES:
            whisper_language = "en"

        text_language = detect_text_language(
            text
        )

        script_language = detect_script_language(
            text
        )

        final_language = "en"

        if script_language == "bn":

            final_language = "bn"

        elif script_language == "hi":

            if (
                whisper_language == "ne"
                and whisper_probability >= 0.45
            ):

                final_language = "ne"

            else:

                final_language = "hi"

        elif script_language in {

            "pa",
            "gu",
            "or",
            "ta",
            "te",
            "kn",
            "ml",

        }:

            final_language = script_language

        elif script_language == "ar":

            if (
                whisper_language == "ur"
                and whisper_probability >= 0.40
            ):

                final_language = "ur"

            elif text_language == "ur":

                final_language = "ur"

            else:

                final_language = "ar"

        elif script_language == "ja":

            final_language = "ja"

        elif script_language == "ko":

            final_language = "ko"

        elif script_language == "zh":

            final_language = "zh"

        elif script_language == "el":

            final_language = "el"

        elif script_language == "he":

            final_language = "he"

        elif script_language == "ru":

            cyrillic_languages = {
                "ru",
                "uk",
                "bg",
                "sr",
            }

            if (
                whisper_language
                in cyrillic_languages
                and whisper_probability >= 0.45
            ):

                final_language = (
                    whisper_language
                )

            else:

                final_language = "ru"

        elif script_language == "th":

            final_language = "th"

        else:

            if text_language != "en":

                if whisper_language == text_language:

                    final_language = (
                        text_language
                    )

                elif (
                    whisper_language != "en"
                    and whisper_probability >= 0.70
                ):

                    final_language = (
                        whisper_language
                    )

                else:

                    final_language = (
                        text_language
                    )

            elif (
                whisper_language != "en"
                and whisper_probability >= 0.45
            ):

                final_language = (
                    whisper_language
                )

            else:

                final_language = "en"

        if final_language not in LANGUAGE_NAMES:

            final_language = "en"

        print("=" * 70)
        print("VOICE LANGUAGE DETECTION")
        print("=" * 70)

        print(
            f"Whisper language    : "
            f"{whisper_language}"
        )

        print(
            f"Whisper probability : "
            f"{whisper_probability:.3f}"
        )

        print(
            f"Transcript          : "
            f"{text}"
        )

        print(
            f"Text language       : "
            f"{text_language}"
        )

        print(
            f"Script language     : "
            f"{script_language}"
        )

        print(
            f"FINAL LANGUAGE      : "
            f"{final_language}"
        )

        print("=" * 70)

        if return_segments:

            return (
                final_language,
                text,
                segment_data
            )

        return (
            final_language,
            text
        )

    except Exception as e:

        raise RuntimeError(
            "Whisper transcription failed: "
            f"{repr(e)}"
        )


# ============================================================
# FILE HASH
# ============================================================

def get_file_hash(
    file_bytes
):

    return hashlib.sha256(
        file_bytes
    ).hexdigest()


# ============================================================
# PDF TEXT + PAGE INFORMATION
# ============================================================

def extract_pdf_data(
    file_bytes
):

    pdf_document = pymupdf.open(
        stream=file_bytes,
        filetype="pdf"
    )

    pages = []

    for index in range(
        len(pdf_document)
    ):

        page = pdf_document[index]

        text = page.get_text(
            "text"
        ).strip()

        pages.append({

            "page_number": index + 1,

            "text": text,
        })

    pdf_document.close()

    return pages


# ============================================================
# PDF PAGE SCREENSHOT
# ============================================================

def render_pdf_page(
    file_bytes,
    page_number
):

    pdf_document = pymupdf.open(
        stream=file_bytes,
        filetype="pdf"
    )

    index = page_number - 1

    if (
        index < 0
        or index >= len(pdf_document)
    ):

        pdf_document.close()

        return None

    page = pdf_document[index]

    matrix = pymupdf.Matrix(
        1.6,
        1.6
    )

    pixmap = page.get_pixmap(
        matrix=matrix,
        alpha=False
    )

    image_bytes = pixmap.tobytes(
        "png"
    )

    pdf_document.close()

    return image_bytes


# ============================================================
# DOCX TEXT
# ============================================================

def extract_text_from_docx(
    file_bytes
):

    with tempfile.NamedTemporaryFile(
        delete=False,
        suffix=".docx"
    ) as tmp:

        tmp.write(
            file_bytes
        )

        tmp_path = tmp.name

    try:

        document = Document(
            tmp_path
        )

        parts = []

        for paragraph in document.paragraphs:

            if paragraph.text.strip():

                parts.append(
                    paragraph.text
                )

        for table in document.tables:

            for row in table.rows:

                row_text = " | ".join(

                    cell.text.strip()

                    for cell in row.cells

                    if cell.text.strip()
                )

                if row_text:

                    parts.append(
                        row_text
                    )

        return "\n".join(
            parts
        ).strip()

    finally:

        try:
            os.remove(
                tmp_path
            )

        except Exception:
            pass


# ============================================================
# CONVERT OFFICE FILE TO PDF
# ============================================================

def convert_office_to_pdf(
    file_bytes,
    extension
):

    soffice = (
        shutil.which("soffice")
        or
        shutil.which("libreoffice")
    )

    if not soffice:

        print(
            "LibreOffice/soffice not found."
        )

        return None

    source_path = os.path.join(

        tempfile.gettempdir(),

        f"office_{uuid.uuid4().hex}"
        f"{extension}"
    )

    output_dir = os.path.join(

        tempfile.gettempdir(),

        f"office_pdf_{uuid.uuid4().hex}"
    )

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    try:

        with open(
            source_path,
            "wb"
        ) as file:

            file.write(
                file_bytes
            )

        command = [

            soffice,

            "--headless",

            "--convert-to",
            "pdf",

            "--outdir",
            output_dir,

            source_path,
        ]

        result = subprocess.run(

            command,

            stdout=subprocess.PIPE,

            stderr=subprocess.PIPE,

            text=True,

            timeout=120,
        )

        print(
            "LibreOffice stdout:",
            result.stdout
        )

        print(
            "LibreOffice stderr:",
            result.stderr
        )

        if result.returncode != 0:

            return None

        pdf_name = (
            Path(source_path).stem
            + ".pdf"
        )

        pdf_path = os.path.join(
            output_dir,
            pdf_name
        )

        if not os.path.exists(
            pdf_path
        ):

            return None

        with open(
            pdf_path,
            "rb"
        ) as pdf_file:

            return pdf_file.read()

    except Exception as e:

        print(
            "Office conversion error:",
            repr(e)
        )

        return None

    finally:

        try:

            if os.path.exists(
                source_path
            ):

                os.remove(
                    source_path
                )

        except Exception:
            pass

        try:

            if os.path.exists(
                output_dir
            ):

                shutil.rmtree(
                    output_dir,
                    ignore_errors=True
                )

        except Exception:
            pass


# ============================================================
# GENERIC PDF DATA
# ============================================================

def extract_pages_from_pdf_bytes(
    file_bytes
):

    try:

        document = pymupdf.open(
            stream=file_bytes,
            filetype="pdf"
        )

        pages = []

        for index in range(
            len(document)
        ):

            page = document[index]

            pages.append({

                "page_number": index + 1,

                "text": page.get_text(
                    "text"
                ).strip(),
            })

        document.close()

        return pages

    except Exception as e:

        print(
            "PDF page extraction error:",
            repr(e)
        )

        return []


# ============================================================
# VIDEO TEMP FILE
# ============================================================

def save_video_to_temp(
    file_bytes,
    extension
):

    path = os.path.join(

        tempfile.gettempdir(),

        f"video_{uuid.uuid4().hex}"
        f"{extension}"
    )

    with open(
        path,
        "wb"
    ) as video:

        video.write(
            file_bytes
        )

    return path


# ============================================================
# EXTRACT VIDEO FRAME
# ============================================================

def extract_video_frame(
    video_path,
    timestamp
):

    if not check_ffmpeg():

        return None

    output_path = os.path.join(

        tempfile.gettempdir(),

        f"frame_{uuid.uuid4().hex}.jpg"
    )

    command = [

        "ffmpeg",

        "-y",

        "-ss",
        str(
            max(
                0,
                timestamp
            )
        ),

        "-i",
        video_path,

        "-frames:v",
        "1",

        "-q:v",
        "2",

        output_path,
    ]

    try:

        result = subprocess.run(

            command,

            stdout=subprocess.PIPE,

            stderr=subprocess.PIPE,

            text=True,

            timeout=30,
        )

        if (
            result.returncode != 0
            or
            not os.path.exists(
                output_path
            )
        ):

            return None

        with open(
            output_path,
            "rb"
        ) as frame:

            data = frame.read()

        return data

    except Exception:

        return None

    finally:

        try:

            if os.path.exists(
                output_path
            ):

                os.remove(
                    output_path
                )

        except Exception:
            pass


# ============================================================
# FORMAT TIMESTAMP
# ============================================================

def format_timestamp(
    seconds
):

    seconds = int(
        max(
            0,
            seconds
        )
    )

    hours = (
        seconds // 3600
    )

    minutes = (
        seconds % 3600
    ) // 60

    secs = (
        seconds % 60
    )

    if hours > 0:

        return (
            f"{hours:02d}:"
            f"{minutes:02d}:"
            f"{secs:02d}"
        )

    return (
        f"{minutes:02d}:"
        f"{secs:02d}"
    )


# ============================================================
# RELEVANCE TOKENIZER
# ============================================================

STOP_WORDS = {

    "the", "is", "are", "was", "were",
    "a", "an", "and", "or", "of", "to",
    "in", "on", "for", "with", "from",
    "this", "that", "what", "why",
    "how", "where", "when", "who",
    "which", "can", "could", "would",
    "please", "tell", "me",

    "ami", "amar", "tumi", "tomar",
    "ki", "kivabe", "kothay", "keno",
    "eta", "ota", "ei", "oi",
    "theke", "niye", "bolo", "dao",

    "mujhe", "mera", "meri", "kya",
    "kaise", "kahan", "kyun",
}


def relevance_tokens(
    text
):

    words = normalize_words(
        text
    )

    return {

        word

        for word in words

        if (
            len(word) >= 2
            and word not in STOP_WORDS
        )
    }


# ============================================================
# RELEVANCE SCORE
# ============================================================

def calculate_relevance(
    question,
    content
):

    if not content:

        return 0

    q_tokens = relevance_tokens(
        question
    )

    c_tokens = relevance_tokens(
        content
    )

    if not q_tokens or not c_tokens:

        return 0

    overlap = len(
        q_tokens.intersection(
            c_tokens
        )
    )

    phrase_bonus = 0

    question_normalized = " ".join(
        normalize_words(
            question
        )
    )

    content_normalized = " ".join(
        normalize_words(
            content
        )
    )

    if (
        len(question_normalized) > 8
        and
        question_normalized
        in content_normalized
    ):

        phrase_bonus = 5

    return (
        overlap * 2
        +
        phrase_bonus
    )


# ============================================================
# FIND RELEVANT PDF PAGES
# ============================================================

def find_relevant_pdf_pages(
    question,
    pages,
    max_pages=2
):

    if not pages:

        return []

    scored = []

    for page in pages:

        score = calculate_relevance(

            question,

            page.get(
                "text",
                ""
            )
        )

        scored.append(
            (
                score,
                page
            )
        )

    scored.sort(
        key=lambda x: x[0],
        reverse=True
    )

    selected = []

    for score, page in scored:

        if score <= 0:

            continue

        selected.append(
            page
        )

        if len(selected) >= max_pages:

            break

    # IMPORTANT:
    # Even if relevance score is zero,
    # still return the best available pages.

    if not selected:

        selected = [

            page

            for _, page in scored[
                :max_pages
            ]
        ]

    return selected


# ============================================================
# WEBSITE TEXT EXTRACTION
# ============================================================

# ============================================================
# GENERIC WEBSITE ACCESS + LOGIN + CRAWLING
# ============================================================

AUTH_KEYWORDS = {
    "login", "log in", "signin", "sign in", "authenticate",
    "authentication", "account", "username", "email", "password",
}

SECURITY_CHALLENGE_KEYWORDS = {
    "captcha", "recaptcha", "verification code", "verify code",
    "one-time password", "otp", "two-factor", "2fa", "authenticator",
    "security code", "passkey",
}


def normalize_url(url, base_url=None):
    url = (url or "").strip()
    if not url:
        return ""
    if base_url:
        url = urljoin(base_url, url)
    url, _ = urldefrag(url)
    if not url.startswith(("http://", "https://")):
        url = "https://" + url
    parsed = urlparse(url)
    return parsed._replace(path=parsed.path or "/", fragment="").geturl()


def same_domain(url1, url2):
    try:
        d1 = urlparse(url1).netloc.lower().split(":")[0]
        d2 = urlparse(url2).netloc.lower().split(":")[0]
        if d1.startswith("www."):
            d1 = d1[4:]
        if d2.startswith("www."):
            d2 = d2[4:]
        return d1 == d2
    except Exception:
        return False


def is_crawlable_url(url):
    if not url:
        return False
    lower = url.lower()
    if lower.startswith(("mailto:", "tel:", "javascript:", "data:", "whatsapp:")):
        return False
    blocked = (
        ".jpg", ".jpeg", ".png", ".gif", ".webp", ".svg", ".ico",
        ".mp4", ".mp3", ".wav", ".avi", ".mov", ".zip", ".rar",
        ".7z", ".exe", ".dmg", ".pdf", ".doc", ".docx", ".xls",
        ".xlsx", ".ppt", ".pptx",
    )
    return not urlparse(lower).path.endswith(blocked)


def clean_website_text(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def _visible_locator(page, selectors):
    for selector in selectors:
        try:
            loc = page.locator(selector)
            for i in range(min(loc.count(), 10)):
                item = loc.nth(i)
                if item.is_visible() and item.is_enabled():
                    return item
        except Exception:
            continue
    return None


def detect_security_challenge(page):
    try:
        text = page.locator("body").inner_text(timeout=5000).lower()
    except Exception:
        text = ""
    return [k for k in SECURITY_CHALLENGE_KEYWORDS if k in text]


def inspect_website_access(url):
    if not PLAYWRIGHT_AVAILABLE:
        return {"success": False, "requires_login": False,
                "error": "Playwright is unavailable. Run: playwright install chromium"}
    url = normalize_url(url)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(
                viewport={"width": 1440, "height": 1000},
                user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                            "AppleWebKit/537.36 (KHTML, like Gecko) "
                            "Chrome/138.0 Safari/537.36"),
            )
            page = context.new_page()
            response = page.goto(url, wait_until="domcontentloaded", timeout=30000)
            try:
                page.wait_for_load_state("networkidle", timeout=10000)
            except Exception:
                pass
            page.wait_for_timeout(1200)
            final_url = normalize_url(page.url)
            try:
                title = page.title()
            except Exception:
                title = ""
            try:
                body = page.locator("body").inner_text(timeout=5000)
            except Exception:
                body = ""
            lower_url = final_url.lower()
            lower = (lower_url + " " + title.lower() + " " + body[:12000].lower())
            password_visible = False
            try:
                pw = page.locator("input[type='password']")
                for i in range(min(pw.count(), 10)):
                    if pw.nth(i).is_visible():
                        password_visible = True
                        break
            except Exception:
                pass
            login_url = None
            try:
                anchors = page.locator("a[href]").evaluate_all(
                    "els => els.map(a => ({href:a.href,text:(a.innerText||'').trim()}))"
                )
                for item in anchors:
                    href = item.get("href", "")
                    text = item.get("text", "").lower()
                    combined = (href.lower() + " " + text)
                    if any(k in combined for k in AUTH_KEYWORDS):
                        candidate = normalize_url(href, final_url)
                        if candidate and same_domain(candidate, final_url):
                            login_url = candidate
                            break
            except Exception:
                pass
            status_login = False
            try:
                status_login = bool(response and response.status in (401, 403))
            except Exception:
                pass
            redirected = any(x in lower_url for x in ("/login", "/signin", "/sign-in", "/authenticate"))
            auth_context = any(k in lower for k in AUTH_KEYWORDS)
            requires_login = status_login or redirected or (password_visible and auth_context)
            if requires_login and not login_url:
                login_url = final_url
            browser.close()
            return {"success": True, "requires_login": requires_login,
                    "login_url": login_url, "final_url": final_url, "title": title}
    except Exception as e:
        return {"success": False, "requires_login": False, "error": str(e)}


def login_to_website(login_url, login_id, login_password):
    if not PLAYWRIGHT_AVAILABLE:
        return {"success": False, "manual_required": False,
                "message": "Playwright is unavailable."}
    login_url = normalize_url(login_url)
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            context = browser.new_context(
                viewport={"width": 1440, "height": 1000},
                user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                            "AppleWebKit/537.36 (KHTML, like Gecko) "
                            "Chrome/138.0 Safari/537.36"),
            )
            page = context.new_page()
            page.goto(login_url, wait_until="domcontentloaded", timeout=30000)
            try:
                page.wait_for_load_state("networkidle", timeout=10000)
            except Exception:
                pass
            page.wait_for_timeout(1000)

            challenge = detect_security_challenge(page)
            if challenge:
                browser.close()
                return {"success": False, "manual_required": True,
                        "message": "This website requires CAPTCHA/OTP/2FA/additional verification. Complete it manually."}

            password_field = _visible_locator(page, [
                "input[type='password']",
                "input[name*='password' i]",
                "input[autocomplete='current-password']",
            ])
            if password_field is None:
                browser.close()
                return {"success": False, "manual_required": False,
                        "message": "Could not find the website password field."}

            username_field = _visible_locator(page, [
                "input[autocomplete='username']",
                "input[type='email']",
                "input[name*='email' i]",
                "input[name*='username' i]",
                "input[name*='user' i]",
                "input[name*='login' i]",
                "input[type='text']",
            ])
            if username_field is None:
                browser.close()
                return {"success": False, "manual_required": False,
                        "message": "Could not find Login ID/Email/Username field."}

            username_field.fill(str(login_id))
            password_field.fill(str(login_password))

            clicked = False
            for selector in [
                "button[type='submit']",
                "input[type='submit']",
                "button",
            ]:
                try:
                    loc = page.locator(selector)
                    for i in range(min(loc.count(), 15)):
                        btn = loc.nth(i)
                        if not btn.is_visible() or not btn.is_enabled():
                            continue
                        ok = btn.evaluate("""
                            el => {
                                const t=(el.innerText||el.value||'').toLowerCase();
                                return el.type==='submit' ||
                                    t.includes('login') || t.includes('log in') ||
                                    t.includes('sign in') || t.includes('signin') ||
                                    t.includes('continue') || t.includes('submit');
                            }
                        """)
                        if ok:
                            btn.click(timeout=10000)
                            clicked = True
                            break
                    if clicked:
                        break
                except Exception:
                    continue

            if not clicked:
                try:
                    password_field.evaluate("""
                        el => { const f=el.closest('form');
                        if(f){ if(f.requestSubmit) f.requestSubmit(); else f.submit(); } }
                    """)
                    clicked = True
                except Exception:
                    pass

            if not clicked:
                browser.close()
                return {"success": False, "manual_required": False,
                        "message": "Could not submit the login form."}

            try:
                page.wait_for_load_state("networkidle", timeout=15000)
            except Exception:
                pass
            page.wait_for_timeout(2500)

            challenge = detect_security_challenge(page)
            if challenge:
                browser.close()
                return {"success": False, "manual_required": True,
                        "message": "The website requested CAPTCHA/OTP/2FA/additional verification."}

            current = page.url.lower()
            try:
                body = page.locator("body").inner_text(timeout=5000).lower()
            except Exception:
                body = ""
            password_still_visible = False
            try:
                pw = page.locator("input[type='password']")
                for i in range(min(pw.count(), 10)):
                    if pw.nth(i).is_visible():
                        password_still_visible = True
                        break
            except Exception:
                pass
            failure_words = [
                "invalid password", "incorrect password", "invalid username",
                "invalid credentials", "wrong password", "login failed",
                "authentication failed", "incorrect login",
            ]
            failed_text = any(x in body for x in failure_words)
            still_login = any(x in current for x in ("/login", "/signin", "/sign-in", "/authenticate"))
            if failed_text or (password_still_visible and still_login):
                browser.close()
                return {"success": False, "manual_required": False,
                        "message": "Login failed. Please check Login ID and Password."}

            state = context.storage_state()
            final_url = page.url
            browser.close()
            return {"success": True, "manual_required": False,
                    "storage_state": state, "final_url": final_url}
    except Exception as e:
        return {"success": False, "manual_required": False,
                "message": f"Login process failed: {e}"}


def _website_page_from_html(url, response):
    soup = BeautifulSoup(response.text, "html.parser")
    for tag in soup(["script", "style", "noscript", "svg", "iframe", "template"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    elements=[]
    for tag in soup.find_all(["h1","h2","h3","h4","h5","h6","p","li","article","section","td","th","button","a"]):
        text=clean_website_text(tag.get_text(" ", strip=True))
        if text:
            elements.append({"tag":tag.name,"text":text[:1500]})
    links=[]
    for a in soup.find_all("a", href=True):
        candidate=normalize_url(a.get("href"), response.url)
        if candidate and same_domain(candidate,url) and is_crawlable_url(candidate):
            links.append(candidate)
    return {
        "url": normalize_url(response.url),
        "title": clean_website_text(title),
        "text": clean_website_text(soup.get_text(" ", strip=True))[:50000],
        "elements": elements[:500],
        "links": list(dict.fromkeys(links))[:100],
        "requires_login": False,
    }


def fetch_website(url):
    url=normalize_url(url)
    if not url:
        raise RuntimeError("Website URL is empty.")
    headers={"User-Agent":("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                             "AppleWebKit/537.36 (KHTML, like Gecko) "
                             "Chrome/138.0 Safari/537.36")}
    session=requests.Session(); session.headers.update(headers)
    queue=[url]; visited=set(); pages=[]
    while queue and len(pages)<25:
        current=queue.pop(0)
        if current in visited: continue
        visited.add(current)
        try:
            response=session.get(current,timeout=20,allow_redirects=True)
            final=normalize_url(response.url)
            ct=response.headers.get("content-type","").lower()
            if response.status_code in (401,403):
                pages.append({"url":final,"title":"","text":"","elements":[],"links":[],"requires_login":True})
                continue
            if "text/html" not in ct and not response.text.lstrip().startswith("<"):
                continue
            page=_website_page_from_html(url,response)
            pages.append(page)
            for link in page["links"]:
                if link not in visited:
                    queue.append(link)
        except Exception:
            continue
    text="\n\n".join(f"PAGE: {x['title']}\nURL: {x['url']}\n{x['text']}" for x in pages)
    return {"url":url,"title":pages[0]["title"] if pages else "","text":text[:120000],"pages":pages,"authenticated":False}


def crawl_authenticated_website(start_url, storage_state, max_pages=25):
    start_url=normalize_url(start_url)
    queue=[start_url]; visited=set(); pages=[]
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=True)
        context=p.chromium.new_context(
            storage_state=storage_state,
            viewport={"width":1440,"height":1000},
            user_agent=("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0 Safari/537.36")
        )
        page=context.new_page()
        while queue and len(pages)<max_pages:
            current=queue.pop(0)
            if current in visited: continue
            visited.add(current)
            try:
                page.goto(current,wait_until="domcontentloaded",timeout=30000)
                try: page.wait_for_load_state("networkidle",timeout=10000)
                except Exception: pass
                page.wait_for_timeout(600)
                final=normalize_url(page.url)
                if not same_domain(final,start_url): continue
                try: title=page.title()
                except Exception: title=""
                try: body=clean_website_text(page.locator("body").inner_text(timeout=5000))
                except Exception: body=""
                requires=any(x in final.lower() for x in ("/login","/signin","/sign-in","/authenticate"))
                try:
                    elements=page.locator("h1,h2,h3,h4,h5,h6,p,li,article,section,td,th,button,a").evaluate_all(
                        "els=>els.map(el=>({tag:el.tagName.toLowerCase(),text:(el.innerText||'').trim()})).filter(x=>x.text).slice(0,500)"
                    )
                except Exception: elements=[]
                links=[]
                try:
                    raw=page.locator("a[href]").evaluate_all("els=>els.map(a=>a.href)")
                    for link in raw:
                        c=normalize_url(link,final)
                        if c and same_domain(c,start_url) and is_crawlable_url(c): links.append(c)
                except Exception: pass
                links=list(dict.fromkeys(links))[:100]
                pages.append({"url":final,"title":title,"text":body[:50000],"elements":elements,
                              "links":links,"requires_login":requires})
                for link in links:
                    if link not in visited: queue.append(link)
            except Exception: continue
        browser.close()
    text="\n\n".join(f"PAGE: {x['title']}\nURL: {x['url']}\n{x['text']}" for x in pages)
    return {"url":start_url,"title":pages[0]["title"] if pages else "","text":text[:120000],
            "pages":pages,"authenticated":True,"storage_state":storage_state}


def find_relevant_website_section(question, sections):
    if not sections: return None
    scored=[]
    for section in sections:
        scored.append((calculate_relevance(question,section.get("text",section.get("section_text",""))),section))
    scored.sort(key=lambda x:x[0],reverse=True)
    return scored[0][1] if scored else None


def find_relevant_website_page(question, website_data):
    pages=website_data.get("pages",[]) if website_data else []
    if not pages: return None
    scored=[]
    for page in pages:
        content=" ".join([page.get("title",""),page.get("url",""),page.get("text","")])
        score=calculate_relevance(question,content)
        for token in relevance_tokens(question):
            if len(token)>=4 and token in (page.get("title","")+" "+page.get("url","")).lower(): score+=5
        scored.append((score,page))
    scored.sort(key=lambda x:x[0],reverse=True)
    return scored[0][1]


def capture_website_sections(url, storage_state=None, question=None):
    if not PLAYWRIGHT_AVAILABLE: return []
    result=[]
    try:
        with sync_playwright() as p:
            browser=p.chromium.launch(headless=True)
            args={"viewport":{"width":1440,"height":1000},"device_scale_factor":1,
                  "user_agent":("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                                 "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/138.0 Safari/537.36")}
            if storage_state: args["storage_state"]=storage_state
            context=browser.new_context(**args)
            page=context.new_page()
            page.goto(url,wait_until="domcontentloaded",timeout=30000)
            try: page.wait_for_load_state("networkidle",timeout=10000)
            except Exception: pass
            page.wait_for_timeout(1000)
            current=page.url.lower()
            if any(x in current for x in ("/login","/signin","/sign-in","/authenticate")):
                context.close(); browser.close(); return []
            for _ in range(8):
                try:
                    page.mouse.wheel(0,1200); page.wait_for_timeout(250)
                except Exception: break
            page.evaluate("window.scrollTo(0,0)"); page.wait_for_timeout(400)
            loc=page.locator("h1,h2,h3,h4,h5,h6,article,section,p,li,td,th,a,button,div")
            q=question or ""
            best=[]
            for i in range(min(loc.count(),2500)):
                try:
                    el=loc.nth(i)
                    if not el.is_visible(): continue
                    text=clean_website_text(el.inner_text(timeout=800))
                    if not text or len(text)>12000: continue
                    box=el.bounding_box()
                    if not box or box["width"]<30 or box["height"]<15: continue
                    score=calculate_relevance(q,text) if q else 0
                    tag=el.evaluate("el=>el.tagName.toLowerCase()")
                    if tag.startswith("h"): score+=5
                    if tag in ("article","section"): score+=3
                    area=box["width"]*box["height"]
                    if area>1440*1000*.8: score-=15
                    elif area>1440*1000*.5: score-=8
                    best.append((score,el,text,tag))
                except Exception: continue
            best.sort(key=lambda x:x[0],reverse=True)
            for score,el,text,tag in best[:12]:
                try:
                    target=el
                    if tag in ("h1","h2","h3","h4","h5","h6","p","a","button"):
                        parent=el.locator("xpath=ancestor::*[self::article or self::section or self::li][1]")
                        if parent.count()>0 and parent.first.is_visible():
                            pt=clean_website_text(parent.first.inner_text(timeout=1000))
                            if pt and len(pt)<=12000 and calculate_relevance(q,pt)>=score-5:
                                target=parent.first; text=pt
                    target.scroll_into_view_if_needed(timeout=5000); page.wait_for_timeout(300)
                    image=target.screenshot(type="png",animations="disabled",timeout=10000)
                    result.append({"text":text,"image":image,"mime_type":"image/png","url":page.url})
                    if len(result)>=5: break
                except Exception: continue
            context.close(); browser.close()
    except Exception as e:
        print("Website screenshot error:",repr(e))
    return result


def capture_exact_website_section(page_url, question, storage_state=None):
    sections=capture_website_sections(page_url,storage_state=storage_state,question=question)
    if not sections:
        inspection=inspect_website_access(page_url)
        if inspection.get("requires_login"):
            return {"success":False,"requires_login":True,"login_url":inspection.get("login_url",page_url)}
        return {"success":False,"requires_login":False,"error":"Could not capture relevant website section."}
    section=find_relevant_website_section(question,sections)
    if not section:
        return {"success":False,"requires_login":False,"error":"No relevant section found."}
    return {"success":True,"requires_login":False,"image":section["image"],"url":section.get("url",page_url),
            "section_text":section.get("text","")}


def get_website_question_source(question, source):
    website_data=source.get("website_data") if source else None
    page=find_relevant_website_page(question,website_data)
    if not page: return [], ""
    state=source.get("storage_state") or website_data.get("storage_state")
    shot=capture_exact_website_section(page.get("url"),question,state)
    if shot.get("requires_login"):
        return [{"type":"website_login_required","login_url":shot.get("login_url",page.get("url"))}], ""
    if not shot.get("success"):
        return [], f"Relevant website page:\n{page.get('url')}\n\nPage title:\n{page.get('title')}\n\nContent:\n{page.get('text','')[:15000]}"
    return [{"type":"image","data":shot["image"],"caption":"🌐 Exact relevant website section",
             "mime_type":"image/png","url":shot["url"],"section_text":shot["section_text"]}], \
           f"Relevant website page: {shot['url']}\n\nExact selected section:\n{shot['section_text']}"



# ============================================================
# PROCESS UPLOADED FILE
# ============================================================

def process_uploaded_file(
    uploaded_file
):

    if uploaded_file is None:

        return {
            "source_type": None
        }

    filename = uploaded_file.name

    extension = Path(
        filename
    ).suffix.lower()

    file_bytes = uploaded_file.getvalue()

    file_hash = get_file_hash(
        file_bytes
    )

    result = {

        "source_type":
        "unknown",

        "filename":
        filename,

        "extension":
        extension,

        "file_hash":
        file_hash,

        "file_bytes":
        file_bytes,

        "uploaded_context":
        "",

        "image_part":
        None,

        "file_part":
        None,

        "pdf_pages":
        [],

        "video_path":
        None,

        "video_segments":
        [],

        "video_language":
        "en",
    }

    # ========================================================
    # PDF
    # ========================================================

    if extension == ".pdf":

        result["source_type"] = "pdf"

        pdf_pages = (
            extract_pdf_data(
                file_bytes
            )
        )

        result["pdf_pages"] = (
            pdf_pages
        )

        result["uploaded_context"] = (
            "\n\n".join(

                page["text"]

                for page in pdf_pages

                if page["text"]
            )
        )

        result["file_part"] = (
            types.Part.from_bytes(

                data=file_bytes,

                mime_type="application/pdf"
            )
        )

        return result

    # ========================================================
    # DOCX
    # ========================================================

    if extension == ".docx":

        result["source_type"] = "docx"

        result["uploaded_context"] = (
            extract_text_from_docx(
                file_bytes
            )
        )

        result["file_part"] = (
            types.Part.from_bytes(

                data=file_bytes,

                mime_type=(
                    "application/vnd.openxmlformats-"
                    "officedocument.wordprocessingml.document"
                )
            )
        )

        converted_pdf = (
            convert_office_to_pdf(

                file_bytes,

                ".docx"
            )
        )

        if converted_pdf:

            result["pdf_pages"] = (
                extract_pages_from_pdf_bytes(
                    converted_pdf
                )
            )

            result["render_pdf_bytes"] = (
                converted_pdf
            )

        return result

    # ========================================================
    # PPTX
    # ========================================================

    if extension == ".pptx":

        result["source_type"] = "pptx"

        try:

            from pptx import Presentation

            with tempfile.NamedTemporaryFile(

                delete=False,

                suffix=".pptx"

            ) as tmp:

                tmp.write(
                    file_bytes
                )

                pptx_path = tmp.name

            try:

                presentation = (
                    Presentation(
                        pptx_path
                    )
                )

                slide_texts = []

                for index, slide in enumerate(
                    presentation.slides
                ):

                    parts = []

                    for shape in slide.shapes:

                        if hasattr(
                            shape,
                            "text"
                        ):

                            if shape.text.strip():

                                parts.append(
                                    shape.text.strip()
                                )

                    slide_texts.append({

                        "page_number":
                        index + 1,

                        "text":
                        "\n".join(parts),
                    })

                result["pdf_pages"] = (
                    slide_texts
                )

                result["uploaded_context"] = (
                    "\n\n".join(

                        page["text"]

                        for page in slide_texts
                    )
                )

            finally:

                try:

                    os.remove(
                        pptx_path
                    )

                except Exception:
                    pass

        except Exception as e:

            print(
                "PPTX extraction error:",
                repr(e)
            )

        result["file_part"] = (
            types.Part.from_bytes(

                data=file_bytes,

                mime_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "presentationml.presentation"
                )
            )
        )

        converted_pdf = (
            convert_office_to_pdf(

                file_bytes,

                ".pptx"
            )
        )

        if converted_pdf:

            result["render_pdf_bytes"] = (
                converted_pdf
            )

            # Use converted PDF pages for
            # actual visual page/slide matching.

            converted_pages = (
                extract_pages_from_pdf_bytes(
                    converted_pdf
                )
            )

            if converted_pages:

                result["pdf_pages"] = (
                    converted_pages
                )

        return result

    # ========================================================
    # TEXT FILES
    # ========================================================

    if extension in {

        ".txt",
        ".md",
        ".csv"

    }:

        result["source_type"] = "text"

        result["uploaded_context"] = (
            file_bytes.decode(
                "utf-8",
                errors="ignore"
            ).strip()
        )

        return result

    # ========================================================
    # IMAGES
    # ========================================================

    image_extensions = {

        ".jpg":
        "image/jpeg",

        ".jpeg":
        "image/jpeg",

        ".png":
        "image/png",

        ".webp":
        "image/webp",

        ".gif":
        "image/gif",
    }

    if extension in image_extensions:

        result["source_type"] = "image"

        result["image_part"] = (
            types.Part.from_bytes(

                data=file_bytes,

                mime_type=(
                    image_extensions[
                        extension
                    ]
                )
            )
        )

        return result

    # ========================================================
    # VIDEO
    # ========================================================

    if extension in {

        ".mp4",
        ".mov",
        ".m4v",
        ".avi",
        ".mkv",

    }:

        result["source_type"] = "video"

        if not check_ffmpeg():

            raise RuntimeError(
                "FFmpeg is required for video upload. "
                "Please install FFmpeg and add it to PATH."
            )

        video_path = save_video_to_temp(

            file_bytes,

            extension
        )

        result["video_path"] = (
            video_path
        )

        wav_path = None

        try:

            wav_path = (
                convert_audio_to_wav(
                    video_path
                )
            )

            whisper_model = (
                load_whisper()
            )

            (
                language,
                transcript,
                segments,
            ) = transcribe_audio(

                whisper_model,

                wav_path,

                return_segments=True,
            )

            result["video_language"] = (
                language
            )

            result["video_segments"] = (
                segments
            )

            result["uploaded_context"] = (
                "\n".join(

                    f"[{format_timestamp(item['start'])}] "
                    f"{item['text']}"

                    for item in segments
                )
            )

        finally:

            try:

                if (
                    wav_path
                    and
                    os.path.exists(
                        wav_path
                    )
                ):

                    os.remove(
                        wav_path
                    )

            except Exception:
                pass

        return result

    return result


# ============================================================
# FIND RELEVANT VIDEO SEGMENTS
# ============================================================

def find_relevant_video_segments(
    question,
    segments,
    max_segments=2
):

    if not segments:

        return []

    scored = []

    for segment in segments:

        score = calculate_relevance(

            question,

            segment.get(
                "text",
                ""
            )
        )

        scored.append(
            (
                score,
                segment
            )
        )

    scored.sort(
        key=lambda x: x[0],
        reverse=True
    )

    selected = []

    for score, segment in scored:

        if score <= 0:

            continue

        selected.append(
            segment
        )

        if len(selected) >= max_segments:

            break

    # IMPORTANT:
    # If no exact text match,
    # still return the best transcript segments.

    if not selected:

        selected = [

            segment

            for _, segment in scored[
                :max_segments
            ]
        ]

    return selected


# ============================================================
# LANGUAGE INSTRUCTION
# ============================================================

def get_language_instruction(
    language_code
):

    language_name = LANGUAGE_NAMES.get(
        language_code,
        "English"
    )

    if language_code == "bn":

        return """
CRITICAL LANGUAGE RULE — BENGALI

The current user question is Bengali.

If the user wrote Bengali using Bengali script,
answer in Bengali script.

If the user wrote Bengali using English/Roman letters
(Banglish/Romanized Bengali), answer in
Romanized Bengali/Banglish.

Examples:

"এটা কী?"
→ Bengali script

"eta ki?"
→ Romanized Bengali

"tumi kemon acho?"
→ Romanized Bengali

Do NOT automatically convert Romanized Bengali
into Bengali script.

The uploaded document language MUST NOT determine
the response language.
"""

    if language_code == "hi":

        return """
CRITICAL LANGUAGE RULE — HINDI

The current user question is Hindi.

If the user uses Devanagari,
answer in Devanagari.

If the user uses Romanized Hindi/Hinglish,
answer naturally in Romanized Hindi/Hinglish.

The uploaded document language MUST NOT determine
the response language.
"""

    if language_code == "en":

        return """
CRITICAL LANGUAGE RULE — ENGLISH

The current user question is English.

Answer completely in natural English.

The uploaded document language MUST NOT
change the response language.
"""

    return f"""
CRITICAL LANGUAGE RULE

The current user question is written in
{language_name}.

Answer naturally in {language_name}.

The uploaded source language MUST NOT
change the response language.

Never switch to English unless the user
explicitly asks for English.
"""


# ============================================================
# HISTORY
# ============================================================

def build_history():

    recent = (
        st.session_state.messages[
            -12:
        ]
    )

    history = []

    for message in recent:

        role = message.get(
            "role",
            ""
        )

        content = message.get(
            "content",
            ""
        )

        if not content:

            continue

        if role == "user":

            history.append(
                f"User: {content}"
            )

        elif role == "assistant":

            history.append(
                f"Assistant: {content}"
            )

    return "\n".join(
        history
    )


# ============================================================
# GEMINI REQUEST
# ============================================================

def generate_with_model(
    model_name,
    prompt,
    image_part=None,
    file_part=None,
    extra_parts=None,
):

    contents = []

    if image_part is not None:

        contents.append(
            image_part
        )

    if file_part is not None:

        contents.append(
            file_part
        )

    if extra_parts:

        contents.extend(
            extra_parts
        )

    contents.append(
        prompt
    )

    response = client.models.generate_content(

        model=model_name,

        contents=contents,
    )

    if response is None:

        raise RuntimeError(
            "Gemini returned an empty response."
        )

    text = getattr(
        response,
        "text",
        None
    )

    if not text:

        raise RuntimeError(
            "Gemini returned no text response."
        )

    return text.strip()


# ============================================================
# RETRY CHECK
# ============================================================

def is_retryable_gemini_error(
    error
):

    text = str(
        error
    ).upper()

    retry_codes = [

        "429",
        "500",
        "502",
        "503",
        "504",

        "RESOURCE_EXHAUSTED",
        "UNAVAILABLE",
        "INTERNAL",
        "BAD_GATEWAY",
        "DEADLINE",
        "TIMEOUT",

        "RATE LIMIT",
        "TOO MANY REQUESTS",
    ]

    return any(
        code in text
        for code in retry_codes
    )


# ============================================================
# GEMINI
# ============================================================

def ask_gemini(
    prompt,
    image_part=None,
    file_part=None,
    extra_parts=None,
):
    """
    Ask Gemini with retry + fallback support.

    The primary model is GEMINI_MODEL (normally gemini-3.5-flash).
    Temporary 429/5xx/unavailable errors are retried briefly and then
    the configured fallback model(s) are tried.

    Existing multimodal support is preserved through image_part,
    file_part and extra_parts.
    """

    models = []

    if GEMINI_MODEL:
        models.append(GEMINI_MODEL)

    for fallback in FALLBACK_MODELS:
        if fallback and fallback not in models:
            models.append(fallback)

    if not models:
        raise RuntimeError(
            "No Gemini model has been configured."
        )

    last_error = None

    for model_index, model_name in enumerate(models):

        # Try each model at most twice. This lets the fallback model
        # be reached quickly when the primary model is temporarily busy.
        max_attempts = 2

        for attempt in range(max_attempts):

            try:
                print(
                    f"Gemini request | "
                    f"model={model_name} | "
                    f"attempt={attempt + 1}"
                )

                result = generate_with_model(
                    model_name=model_name,
                    prompt=prompt,
                    image_part=image_part,
                    file_part=file_part,
                    extra_parts=extra_parts,
                )

                if result:
                    return result

                raise RuntimeError(
                    "Gemini returned an empty response."
                )

            except Exception as e:
                last_error = e

                print(
                    f"Gemini error | "
                    f"model={model_name} | "
                    f"attempt={attempt + 1} | "
                    f"error={repr(e)}"
                )

                if not is_retryable_gemini_error(e):
                    raise RuntimeError(
                        f"Gemini API error: {e}"
                    ) from e

                if attempt < max_attempts - 1:
                    delay = min(
                        3 * (2 ** attempt),
                        10,
                    )

                    print(
                        f"Gemini retrying in {delay} seconds..."
                    )
                    time.sleep(delay)

        # If this model was exhausted, continue to the next fallback.
        if model_index < len(models) - 1:
            print(
                f"Gemini model unavailable: {model_name}. "
                f"Trying fallback model: {models[model_index + 1]}"
            )

    raise RuntimeError(
        "Gemini API is temporarily unavailable.\n\n"
        f"Last error: {last_error}"
    )


# ============================================================
# GET SOURCE CONTEXT + VISUAL
# ============================================================


def get_question_source_media(question):
    source=st.session_state.source_data
    if not source:
        return [], ""
    source_type=source.get("source_type")
    media=[]
    extra_context=""

    if source_type == "pdf":
        pages=find_relevant_pdf_pages(question,source.get("pdf_pages",[]),max_pages=2)
        for page in pages:
            image_bytes=render_pdf_page(source["file_bytes"],page["page_number"])
            if image_bytes:
                media.append({"type":"image","data":image_bytes,
                              "caption":f"📄 Relevant PDF page {page['page_number']}","mime_type":"image/png"})

    elif source_type in {"docx","pptx"}:
        pages=find_relevant_pdf_pages(question,source.get("pdf_pages",[]),max_pages=2)
        render_pdf_bytes=source.get("render_pdf_bytes")
        if render_pdf_bytes:
            for page in pages:
                image_bytes=render_pdf_page(render_pdf_bytes,page["page_number"])
                if image_bytes:
                    label="DOCX page" if source_type=="docx" else "PPTX slide"
                    media.append({"type":"image","data":image_bytes,
                                  "caption":f"📄 Relevant {label} {page['page_number']}","mime_type":"image/png"})

    elif source_type == "video":
        segments=find_relevant_video_segments(question,source.get("video_segments",[]),max_segments=2)
        video_path=source.get("video_path")
        if video_path and os.path.exists(video_path):
            for segment in segments:
                timestamp=segment["start"]+(segment["end"]-segment["start"])/2
                frame=extract_video_frame(video_path,timestamp)
                if frame:
                    media.append({"type":"image","data":frame,
                                  "caption":f"🎬 Relevant video frame at {format_timestamp(timestamp)}",
                                  "timestamp":timestamp,"video_text":segment["text"],"mime_type":"image/jpeg"})

    elif source_type == "website":
        website_media, website_context=get_website_question_source(question,source)
        media.extend(website_media)
        extra_context=website_context

    elif source_type == "image":
        image_data=source.get("file_bytes",b"")
        if image_data:
            image_extensions={".jpg":"image/jpeg",".jpeg":"image/jpeg",".png":"image/png",
                              ".webp":"image/webp",".gif":"image/gif"}
            mime=image_extensions.get(source.get("extension",".png").lower(),"image/png")
            media.append({"type":"image","data":image_data,"caption":"🖼️ Uploaded image",
                          "mime_type":mime,"already_sent":True})

    return media, extra_context



# ============================================================
# ASK AI
# ============================================================

def ask_ai(
    user_message,
    language_code,
    relevant_media=None,
):

    if language_code not in LANGUAGE_NAMES:

        language_code = "en"

    # --------------------------------------------------------
    # ALWAYS DETECT LANGUAGE FROM CURRENT QUESTION
    # --------------------------------------------------------

    detected_question_language = (
        detect_text_language(
            user_message
        )
    )

    if detected_question_language in LANGUAGE_NAMES:

        language_code = (
            detected_question_language
        )

    language_name = LANGUAGE_NAMES.get(
        language_code,
        "English"
    )

    language_instruction = (
        get_language_instruction(
            language_code
        )
    )

    history = build_history()

    source = (
        st.session_state.source_data
    )

    uploaded_context = ""

    image_part = None

    file_part = None

    extra_parts = []

    source_instruction = ""

    # ========================================================
    # ADD RELEVANT VISUALS TO GEMINI
    # ========================================================

    if relevant_media:

        for media in relevant_media:

            # Uploaded image is already
            # supplied using image_part.
            if media.get(
                "already_sent"
            ):

                continue

            image_data = media.get(
                "data"
            )

            if image_data:

                try:

                    extra_parts.append(

                        types.Part.from_bytes(

                            data=image_data,

                            mime_type=media.get(
                                "mime_type",
                                "image/png"
                            ),
                        )
                    )

                except Exception as e:

                    print(
                        "Visual context error:",
                        repr(e)
                    )

    # ========================================================
    # SOURCE
    # ========================================================

    if source:

        uploaded_context = source.get(
            "uploaded_context",
            ""
        )

        image_part = source.get(
            "image_part"
        )

        file_part = source.get(
            "file_part"
        )

        source_type = source.get(
            "source_type"
        )

        filename = source.get(
            "filename",
            ""
        )

        source_instruction = f"""

============================================================
CURRENT SOURCE
============================================================

Source type:
{source_type}

Source name:
{filename}

Use the uploaded source only when relevant
to the current user question.

Do not invent information.

The source language MUST NOT determine
the answer language.

The CURRENT USER QUESTION determines
the answer language.

"""

        if uploaded_context:

            source_instruction += f"""

SOURCE TEXT / TRANSCRIPT:

{uploaded_context[:50000]}

============================================================
END SOURCE TEXT
============================================================
"""

        # ----------------------------------------------------
        # VIDEO
        # ----------------------------------------------------

        if source_type == "video":

            source_instruction += """

This is a video source.

The transcript contains timestamps.

When answering a question about the video,
use the transcript to identify the relevant
part of the video.

Relevant video frames may also be provided
as visual context.

If a timestamp is useful, mention the relevant
timestamp naturally.

Do not invent timestamps.
"""

        # ----------------------------------------------------
        # WEBSITE
        # ----------------------------------------------------

        elif source_type == "website":

            source_instruction += """

This is a website source.

Use the extracted website content to answer.

Relevant website screenshots may be provided
as visual context.

Only discuss information actually available
on the website.
"""

        # ----------------------------------------------------
        # DOCUMENT
        # ----------------------------------------------------

        elif source_type in {
            "pdf",
            "docx",
            "pptx"
        }:

            source_instruction += """

This is a document source.

Use the document content to answer.

Relevant document pages or slides may be
provided as visual context.

When visual information is important,
inspect the provided page/slide image.
"""

        # ----------------------------------------------------
        # IMAGE
        # ----------------------------------------------------

        elif source_type == "image":

            source_instruction += """

This is an uploaded image.

Inspect the uploaded image carefully.

Use visual information from the image
when answering the current question.

Do not assume information that cannot
be seen or reasonably inferred from
the image.
"""

    # ========================================================
    # FINAL PROMPT
    # ========================================================

    prompt = f"""
You are a professional multilingual AI voice assistant.

You answer questions using conversation context,
uploaded sources, and relevant visual context.

============================================================
CURRENT USER QUESTION
============================================================

{user_message}

============================================================
LANGUAGE PRIORITY
============================================================

The CURRENT USER QUESTION determines the answer language.

Priority:

1. Current user question
2. Explicit language request
3. Conversation context
4. Uploaded source

The source language MUST NEVER override
the current question language.

============================================================
DETECTED LANGUAGE
============================================================

Code:
{language_code}

Language:
{language_name}

============================================================
LANGUAGE RULE
============================================================

{language_instruction}

============================================================
ABSOLUTE CURRENT-QUESTION LANGUAGE RULE
============================================================

Answer the user's CURRENT QUESTION in the
same language/style used by the user.

Do NOT answer according to the language
of the uploaded document.

Do NOT answer according to the language
of the uploaded PDF.

Do NOT answer according to the language
of the uploaded DOCX.

Do NOT answer according to the language
of the uploaded PPTX.

Do NOT answer according to the language
of the uploaded image.

Do NOT answer according to the language
of the uploaded video.

Do NOT answer according to the language
of the website.

The CURRENT USER QUESTION has the highest
language priority.

Examples:

User asks in English + Bengali PDF
→ Answer in English.

User asks in Banglish + English PDF
→ Answer in Banglish.

User asks in Bengali script + English PDF
→ Answer in Bengali script.

User asks in Hinglish + English document
→ Answer in Hinglish.

User asks in Hindi Devanagari + English document
→ Answer in Hindi Devanagari.

============================================================
IMPORTANT
============================================================

If the user writes Roman Bengali/Banglish,
answer in Roman Bengali/Banglish.

If the user writes Bengali script,
answer in Bengali script.

If the user writes Roman Hindi/Hinglish,
answer in Roman Hindi/Hinglish.

If the user writes Hindi Devanagari,
answer in Hindi Devanagari.

If the user writes English,
answer in English.

Do not translate unless requested.

Do not switch language because
the source is in English.

Do not switch language because
the source is in Bengali.

Do not switch language because
the source is in Hindi.

Do not switch language because
the source is in any other language.

============================================================
GENERAL BEHAVIOUR
============================================================

1. Answer the actual question.

2. Use the uploaded source when relevant.

3. Do not invent information.

4. Keep answers reasonably concise.

5. For technical questions,
   explain clearly.

6. If code is requested,
   provide complete code when appropriate.

7. The answer will be spoken aloud,
   so use natural conversational sentences.

8. Do not add unnecessary language labels.

9. Do not explain language detection.

10. Do not provide multiple translations.

11. Do not say "Here is the translation."

12. Directly answer the user.

13. For video questions, use the transcript,
    relevant visual frame and timestamps
    when relevant.

14. For document questions, use relevant
    document content and relevant page/slide
    visual context when available.

15. For website questions, use relevant
    website content and relevant screenshot
    when available.

16. For image questions, inspect the uploaded
    image and answer based on visible content.

17. Never claim that a screenshot/frame/page
    was inspected if no visual was provided.

18. Never invent a timestamp.

============================================================
RECENT CONVERSATION
============================================================

{history}

============================================================
SOURCE INFORMATION
============================================================

{source_instruction}

============================================================
VISUAL CONTEXT
============================================================

Relevant screenshots, document pages,
slides, video frames, or uploaded images
may be attached separately to this request.

Use those visuals when relevant to the
current question.

============================================================
ABSOLUTE LANGUAGE LOCK
============================================================

Answer ONLY in the language/style of the
CURRENT USER QUESTION.

Current language:
{language_name}

Never allow the uploaded document,
PDF, DOCX, PPTX, image, video or website
language to override this.

Return ONLY the actual answer.
"""

    return ask_gemini(

        prompt=prompt,

        image_part=image_part,

        file_part=file_part,

        extra_parts=extra_parts,
    )


# ============================================================
# TEXT TO SPEECH
# ============================================================

def text_to_speech(
    text,
    language_code
):

    if not text or not text.strip():

        print(
            "TTS ERROR: Empty text."
        )

        return None

    if language_code not in TTS_VOICES:

        language_code = "en"

    voice = TTS_VOICES.get(

        language_code,

        TTS_VOICES["en"]
    )

    output_file = os.path.join(

        tempfile.gettempdir(),

        f"tts_{uuid.uuid4().hex}.mp3"
    )

    command = [

        sys.executable,

        "-m",
        "edge_tts",

        "--voice",
        voice,

        "--text",
        text.strip(),

        "--write-media",
        output_file,
    ]

    try:

        result = subprocess.run(

            command,

            stdout=subprocess.PIPE,

            stderr=subprocess.PIPE,

            text=True,

            encoding="utf-8",

            errors="replace",

            timeout=120,
        )

        if result.stdout:

            print(
                "TTS STDOUT:",
                result.stdout
            )

        if result.stderr:

            print(
                "TTS STDERR:",
                result.stderr
            )

        if result.returncode != 0:

            print(
                "TTS ERROR: edge-tts failed."
            )

            return None

        if not os.path.exists(
            output_file
        ):

            return None

        file_size = os.path.getsize(
            output_file
        )

        if file_size < 1000:

            try:

                os.remove(
                    output_file
                )

            except Exception:
                pass

            return None

        return output_file

    except subprocess.TimeoutExpired:

        print(
            "TTS ERROR: timeout."
        )

        return None

    except Exception as e:

        print(
            "TTS ERROR:",
            repr(e)
        )

        return None


# ============================================================
# SESSION STATE
# ============================================================

DEFAULTS = {
    "messages": [],
    "db_session_id": "",
    "db_loaded": False,
    "pending_transcript": "",
    "pending_language": "en",
    "recorder_key": 0,
    "voice_ready": False,
    "source_data": None,
    "source_filename": "",
    "website_url_input": "",
    "website_login_required": False,
    "website_login_url": "",
    "website_original_url": "",
    "website_authenticated": False,
    "pending_website_question": "",
}


for key, value in DEFAULTS.items():

    if key not in st.session_state:

        st.session_state[key] = value


# ============================================================
# SQL DATABASE SESSION
# ============================================================

def init_sql_session():
    """Create/load the persistent SQL chat session once per Streamlit session."""
    if st.session_state.get("db_loaded"):
        return

    try:
        session_id = st.session_state.get("db_session_id") or create_session()
        st.session_state.db_session_id = session_id
        st.session_state.messages = load_messages(session_id)
        st.session_state.db_loaded = True
        print("SQL DATABASE: session loaded", session_id)
    except Exception as e:
        st.session_state.db_loaded = True
        print("SQL DATABASE LOAD ERROR:", repr(e))


# ============================================================
# CLEAR SOURCE
# ============================================================


def clear_source():
    source=st.session_state.get("source_data")
    if source:
        video_path=source.get("video_path")
        if video_path:
            try:
                if os.path.exists(video_path): os.remove(video_path)
            except Exception: pass
    st.session_state.source_data=None
    st.session_state.source_filename=""
    st.session_state.website_login_required=False
    st.session_state.website_login_url=""
    st.session_state.website_original_url=""
    st.session_state.website_authenticated=False
    st.session_state.pending_website_question=""



# ============================================================
# NEW CHAT
# ============================================================

def start_new_chat():

    st.session_state.messages = []

    st.session_state.pending_transcript = ""

    st.session_state.pending_language = "en"

    st.session_state.voice_ready = False

    clear_source()

    try:
        st.session_state.db_session_id = create_session()
        touch_session(st.session_state.db_session_id, "New Chat")
    except Exception as e:
        print("SQL DATABASE NEW SESSION ERROR:", repr(e))

    st.session_state.recorder_key += 1


# ============================================================
# PROCESS USER MESSAGE
# ============================================================


def process_user_message(user_message, language_code):
    if not user_message: return
    user_message=user_message.strip()
    if not user_message: return
    detected=detect_text_language(user_message)
    if detected in LANGUAGE_NAMES: language_code=detected
    if language_code not in LANGUAGE_NAMES: language_code="en"

    relevant_media=[]
    source_context=""
    try:
        relevant_media,source_context=get_question_source_media(user_message)
    except Exception as e:
        print("SOURCE MEDIA ERROR:",repr(e))

    login_items=[m for m in relevant_media if isinstance(m,dict) and m.get("type")=="website_login_required"]
    if login_items:
        st.session_state.pending_website_question=user_message
        st.session_state.website_login_required=True
        st.session_state.website_login_url=login_items[0].get("login_url") or st.session_state.get("website_url_input","")
        st.warning("🔐 This page requires login. Please enter your Login ID/Email/Username and Password above.")
        return

    user_record = {"role": "user", "content": user_message, "language": language_code}
    st.session_state.messages.append(user_record)
    try:
        save_message(
            st.session_state.get("db_session_id"),
            "user",
            user_message,
            language=language_code,
        )
    except Exception as e:
        print("SQL USER MESSAGE SAVE ERROR:", repr(e))
    maybe_update_chat_title(user_message)
    try:
        with st.spinner("🤖 Thinking..."):
            response=ask_ai(user_message=user_message,language_code=language_code,relevant_media=relevant_media)
    except Exception as e:
        st.error("❌ Gemini request failed.")
        print("GEMINI ERROR:",repr(e))
        if st.session_state.messages and st.session_state.messages[-1].get("role")=="user":
            st.session_state.messages.pop()
        st.exception(RuntimeError(str(e)))
        return

    assistant_message={"role":"assistant","content":response,"tts_language":language_code,"source_media":relevant_media}
    for media in relevant_media:
        if media.get("timestamp") is not None:
            assistant_message.setdefault("timestamps",[]).append(media["timestamp"])
    st.session_state.messages.append(assistant_message)
    try:
        save_message(
            st.session_state.get("db_session_id"),
            "assistant",
            response,
            tts_language=language_code,
            media=relevant_media,
        )
    except Exception as e:
        print("SQL ASSISTANT MESSAGE SAVE ERROR:", repr(e))
    with st.spinner("🔊 Generating voice..."):
        audio_file=text_to_speech(text=response,language_code=language_code)
    if audio_file:
        st.session_state.messages[-1]["audio_file"]=audio_file
    else:
        st.warning("⚠️ Text response generated, but voice generation failed. Check the terminal for TTS ERROR.")



# ============================================================
# RENDER MESSAGES
# ============================================================

def render_messages():

    for message in (
        st.session_state.messages
    ):

        role = message.get(
            "role",
            "assistant"
        )

        content = message.get(
            "content",
            ""
        )

        with st.chat_message(
            role
        ):

            st.markdown(
                content
            )

            # ------------------------------------------------
            # SOURCE MEDIA
            # ------------------------------------------------

            source_media = message.get(
                "source_media",
                []
            )

            for media in source_media:

                image_data = media.get(
                    "data"
                )

                if not image_data:

                    continue

                caption = media.get(
                    "caption",
                    "Relevant source"
                )

                timestamp = media.get(
                    "timestamp"
                )

                if timestamp is not None:

                    st.image(

                        image_data,

                        caption=(
                            f"{caption} — "
                            f"{format_timestamp(timestamp)}"
                        ),

                        use_container_width=True,
                    )

                    video_text = media.get(
                        "video_text"
                    )

                    if video_text:

                        st.caption(
                            f"📝 Transcript: "
                            f"{video_text}"
                        )

                else:

                    st.image(

                        image_data,

                        caption=caption,

                        use_container_width=True,
                    )

            # ------------------------------------------------
            # AUDIO
            # ------------------------------------------------

            audio_file = message.get(
                "audio_file"
            )

            if (

                role == "assistant"

                and audio_file

                and os.path.exists(
                    audio_file
                )

            ):

                try:

                    with open(
                        audio_file,
                        "rb"
                    ) as audio:

                        audio_bytes = (
                            audio.read()
                        )

                    st.audio(

                        audio_bytes,

                        format="audio/mp3",

                        autoplay=True,
                    )

                except Exception as e:

                    print(
                        "Audio rendering error:",
                        repr(e)
                    )


# ============================================================
# PERSISTENT CHAT HISTORY
# ============================================================

def _history_title(session):
    title = (session.get("title") or "New Chat").strip()
    if title == "New Chat":
        return "New Chat"
    return title[:42] + ("…" if len(title) > 42 else "")


def load_saved_chat(session_id):
    if not session_id:
        return
    try:
        st.session_state.db_session_id = session_id
        st.session_state.messages = load_messages(session_id)
        st.session_state.db_loaded = True
        clear_source()
        st.session_state.pending_transcript = ""
        st.session_state.voice_ready = False
        st.session_state.recorder_key += 1
        print("SQL DATABASE: history opened", session_id)
    except Exception as e:
        print("SQL HISTORY LOAD ERROR:", repr(e))
        st.error("Could not load this chat history.")


def maybe_update_chat_title(user_message):
    session_id = st.session_state.get("db_session_id")
    if not session_id:
        return
    # Use the first user message as a ChatGPT-style title.
    user_messages = [
        m.get("content", "").strip()
        for m in st.session_state.get("messages", [])
        if m.get("role") == "user" and m.get("content", "").strip()
    ]
    if len(user_messages) != 1:
        return
    title = user_messages[0].replace("\n", " ").strip()
    if len(title) > 48:
        title = title[:48].rstrip() + "…"
    try:
        touch_session(session_id, title or "New Chat")
    except Exception as e:
        print("SQL TITLE UPDATE ERROR:", repr(e))


def render_chat_history():
    try:
        sessions = list_sessions(limit=100)
    except Exception as e:
        print("SQL HISTORY LIST ERROR:", repr(e))
        st.caption("Chat history is temporarily unavailable.")
        return

    if not sessions:
        st.caption("No previous chats yet.")
        return

    current_id = st.session_state.get("db_session_id")
    for session in sessions:
        sid = session.get("id")
        title = _history_title(session)
        is_current = sid == current_id
        label = ("🟢 " if is_current else "💬 ") + title
        if st.button(
            label,
            key=f"history_{sid}",
            use_container_width=True,
            help=f"Open chat from {session.get('updated_at', '')}",
        ):
            load_saved_chat(sid)
            st.rerun()


# ============================================================
# SIDEBAR
# ============================================================

def render_sidebar():

    with st.sidebar:

        st.title(
            "🎙️ AI Voice Assistant"
        )

        st.caption(
            "Multilingual AI voice assistant"
        )

        st.divider()

        # ----------------------------------------------------
        # NEW CHAT
        # ----------------------------------------------------

        if st.button(

            "🆕 New Chat",

            use_container_width=True

        ):

            start_new_chat()

            st.rerun()

        # ----------------------------------------------------
        # CHAT HISTORY
        # ----------------------------------------------------

        st.subheader("💬 Chat History")
        render_chat_history()

        # ----------------------------------------------------
        # CLEAR SOURCE
        # ----------------------------------------------------

        if st.button(

            "🧹 Clear Uploaded Source",

            use_container_width=True

        ):

            clear_source()

            st.rerun()

        st.divider()

        # ----------------------------------------------------
        # SEARCH
        # ----------------------------------------------------

        st.subheader(
            "🔎 Search Chat"
        )

        search_text = st.text_input(

            "Search messages",

            placeholder="Type to search..."
        )

        if search_text:

            search_lower = (
                search_text.lower()
            )

            found = False

            for message in (
                st.session_state.messages
            ):

                content = message.get(
                    "content",
                    ""
                )

                if search_lower in (
                    content.lower()
                ):

                    found = True

                    st.write(
                        content[:300]
                    )

            if not found:

                st.caption(
                    "No matching message found."
                )

        st.divider()

        # ----------------------------------------------------
        # SOURCE STATUS
        # ----------------------------------------------------

        st.subheader(
            "📚 Current Source"
        )

        source = (
            st.session_state.source_data
        )

        if source:

            st.success(

                f"✅ "
                f"{source.get('filename', 'Source')}"
            )

            st.caption(

                f"Type: "
                f"{source.get('source_type')}"
            )

        else:

            st.caption(
                "No source uploaded."
            )

        st.divider()

        # ----------------------------------------------------
        # LANGUAGES
        # ----------------------------------------------------

        st.subheader(
            "🌍 Supported Languages"
        )

        st.caption(

            f"{len(LANGUAGE_NAMES)} "
            "languages supported"
        )

        with st.expander(
            "View languages"
        ):

            for code, name in (
                LANGUAGE_NAMES.items()
            ):

                st.write(
                    f"• {code.upper()} — {name}"
                )

        st.divider()

        # ----------------------------------------------------
        # CLEAR HISTORY
        # ----------------------------------------------------

        if st.button(

            "🗑️ Clear Chat History",

            use_container_width=True

        ):

            st.session_state.messages = []
            try:
                clear_session_messages(st.session_state.get("db_session_id"))
            except Exception as e:
                print("SQL CLEAR HISTORY ERROR:", repr(e))

            st.rerun()

        st.divider()

        st.caption(
            f"Whisper: "
            f"{WHISPER_MODEL_NAME}"
        )

        st.caption(
            f"Gemini: "
            f"{GEMINI_MODEL}"
        )

        st.caption(
            f"Playwright: "
            f"{'Available' if PLAYWRIGHT_AVAILABLE else 'Not available'}"
        )

        st.caption(
            f"FFmpeg: "
            f"{'Available' if check_ffmpeg() else 'Not available'}"
        )


# ============================================================
# MAIN
# ============================================================

def main():

    init_sql_session()

    render_sidebar()

    st.title(
        "🎙️ Multilingual AI Voice Assistant"
    )

    st.markdown(
        "Speak naturally in your language "
        "and get text + voice AI responses."
    )

    st.divider()

    # ========================================================
    # WEBSITE URL + PUBLIC / LOGIN-PROTECTED SUPPORT
    # ========================================================

    st.subheader("🌐 Website Analysis")

    website_url = st.text_input(
        "Enter website URL",
        placeholder="https://example.com",
        key="website_url_input",
    )

    if st.button("🔍 Analyze Website", use_container_width=True):
        if not website_url.strip():
            st.warning("Please enter a website URL.")
        else:
            website_url = normalize_url(website_url)
            try:
                with st.spinner("🌐 Checking website access..."):
                    inspection = inspect_website_access(website_url)

                if not inspection.get("success"):
                    st.error(f"❌ Website access check failed: {inspection.get('error','Unknown error')}")
                elif inspection.get("requires_login"):
                    st.session_state.website_login_required = True
                    st.session_state.website_login_url = inspection.get("login_url") or website_url
                    st.session_state.website_original_url = website_url
                    st.warning("🔐 This website requires login. Enter your credentials below.")
                else:
                    with st.spinner("🌐 Crawling public website pages..."):
                        website_data = fetch_website(website_url)
                    st.session_state.source_data = {
                        "source_type": "website",
                        "filename": website_url,
                        "extension": "",
                        "file_hash": "",
                        "file_bytes": b"",
                        "uploaded_context": website_data.get("text", ""),
                        "image_part": None,
                        "file_part": None,
                        "website_url": website_url,
                        "website_title": website_data.get("title", ""),
                        "website_data": website_data,
                        "website_authenticated": False,
                    }
                    st.session_state.source_filename = website_url
                    st.session_state.website_authenticated = False
                    st.session_state.website_login_required = False
                    st.success(f"✅ Website analyzed successfully. {len(website_data.get('pages', []))} same-domain pages scanned.")
                    if website_data.get("title"):
                        st.info(f"📌 {website_data['title']}")
            except Exception as e:
                st.error(f"❌ Website analysis failed: {e}")

    # --------------------------------------------------------
    # LOGIN FORM
    # --------------------------------------------------------
    if st.session_state.get("website_login_required", False):
        st.markdown("---")
        st.warning("🔐 Login is required to access this website/page.")
        login_url = st.session_state.get("website_login_url", "")
        if login_url:
            st.caption(f"Login page: {login_url}")

        with st.form("website_login_form"):
            login_id = st.text_input("Login ID / Email / Username", placeholder="Enter login ID")
            login_password = st.text_input("Password", type="password", placeholder="Enter password")
            login_submit = st.form_submit_button("🔐 Login & Analyze Website", use_container_width=True)

        if login_submit:
            if not login_id.strip():
                st.error("Please enter Login ID / Email / Username.")
            elif not login_password:
                st.error("Please enter your password.")
            else:
                with st.spinner("🔐 Logging in..."):
                    login_result = login_to_website(login_url, login_id, login_password)
                if login_result.get("manual_required"):
                    st.warning(login_result.get("message", "Manual verification is required."))
                elif not login_result.get("success"):
                    st.error(login_result.get("message", "Login failed."))
                else:
                    state = login_result.get("storage_state")
                    original_url = st.session_state.get("website_original_url") or st.session_state.get("website_url_input")
                    try:
                        with st.spinner("✅ Login successful. Crawling authenticated pages..."):
                            website_data = crawl_authenticated_website(original_url, state, max_pages=25)
                        website_data["storage_state"] = state
                        st.session_state.source_data = {
                            "source_type": "website",
                            "filename": original_url,
                            "extension": "",
                            "file_hash": "",
                            "file_bytes": b"",
                            "uploaded_context": website_data.get("text", ""),
                            "image_part": None,
                            "file_part": None,
                            "website_url": original_url,
                            "website_title": website_data.get("title", ""),
                            "website_data": website_data,
                            "website_authenticated": True,
                            "storage_state": state,
                        }
                        st.session_state.source_filename = original_url
                        st.session_state.website_authenticated = True
                        st.session_state.website_login_required = False
                        st.session_state.website_login_url = ""
                        # Password is intentionally NOT stored.
                        st.success(f"✅ Login successful. {len(website_data.get('pages', []))} authenticated pages scanned.")
                        st.rerun()
                    except Exception as e:
                        st.error(f"❌ Authenticated website analysis failed: {e}")

    # --------------------------------------------------------
    # AUTHENTICATED STATUS
    # --------------------------------------------------------
    if st.session_state.get("website_authenticated", False):
        st.success("🔐 Logged-in website session is active for this Streamlit session.")
        if st.button("🚪 Logout / Clear Website Session", use_container_width=True):
            clear_source()
            st.rerun()

    # --------------------------------------------------------
    # SHOW WEBSITE STATUS
    # --------------------------------------------------------
    current_source = st.session_state.source_data
    if current_source and current_source.get("source_type") == "website":
        with st.expander("🌐 View analyzed website content"):
            website_data = current_source.get("website_data", {})
            website_text = website_data.get("text", current_source.get("uploaded_context", ""))
            if len(website_text) > 8000:
                website_text = website_text[:8000] + "\n\n...[truncated]"
            st.text_area("Website content", website_text, height=250, disabled=True)
            st.caption(f"Pages scanned: {len(website_data.get('pages', []))}")

    st.divider()

    # ========================================================
    # CAMERA CAPTURE
    # ========================================================

    st.subheader("📷 Camera")
    camera_photo = st.camera_input(
        "Take a photo",
        key="camera_capture",
    )

    if camera_photo is not None:
        camera_bytes = camera_photo.getvalue()
        camera_hash = get_file_hash(camera_bytes)
        current_hash = (
            st.session_state.source_data.get("file_hash")
            if st.session_state.get("source_data")
            else None
        )

        if camera_hash != current_hash:
            try:
                camera_source = {
                    "source_type": "image",
                    "filename": f"camera_{camera_hash[:12]}.jpg",
                    "extension": ".jpg",
                    "file_hash": camera_hash,
                    "file_bytes": camera_bytes,
                    "uploaded_context": "",
                    "image_part": types.Part.from_bytes(
                        data=camera_bytes,
                        mime_type="image/jpeg",
                    ),
                    "file_part": None,
                    "camera_capture": True,
                }
                st.session_state.source_data = camera_source
                st.session_state.source_filename = camera_source["filename"]
                save_source(
                    st.session_state.get("db_session_id"),
                    camera_source,
                )
                st.success("📷 Camera photo is ready. Ask a question about it.")
            except Exception as e:
                print("CAMERA PROCESS ERROR:", repr(e))
                st.error(f"❌ Camera image processing failed: {e}")

        st.image(
            camera_bytes,
            caption="📷 Camera photo",
            use_container_width=True,
        )

    # ========================================================
    # FILE UPLOAD
    # ========================================================

    uploaded_file = st.file_uploader(

        "📎 Upload document, image or video",

        type=[

            "pdf",
            "docx",
            "pptx",

            "txt",
            "md",
            "csv",

            "jpg",
            "jpeg",
            "png",
            "webp",
            "gif",

            "mp4",
            "mov",
            "m4v",
            "avi",
            "mkv",
        ],
    )

    if uploaded_file is not None:

        file_hash = get_file_hash(
            uploaded_file.getvalue()
        )

        current_hash = None

        if st.session_state.source_data:

            current_hash = (
                st.session_state.source_data.get(
                    "file_hash"
                )
            )

        if file_hash != current_hash:

            try:

                with st.spinner(
                    "📄 Processing uploaded source..."
                ):

                    source_data = (
                        process_uploaded_file(
                            uploaded_file
                        )
                    )

                    st.session_state.source_data = (
                        source_data
                    )

                    st.session_state.source_filename = (
                        uploaded_file.name
                    )

                    try:
                        source_data["filename"] = uploaded_file.name
                        save_source(
                            st.session_state.get("db_session_id"),
                            source_data,
                        )
                    except Exception as e:
                        print("SQL SOURCE SAVE ERROR:", repr(e))

                st.success(
                    f"✅ "
                    f"{uploaded_file.name} processed."
                )

            except Exception as e:

                st.error(
                    f"❌ File processing failed: {e}"
                )

        source = (
            st.session_state.source_data
        )

        extension = Path(
            uploaded_file.name
        ).suffix.lower()

        # ----------------------------------------------------
        # IMAGE PREVIEW
        # ----------------------------------------------------

        if extension in {

            ".jpg",
            ".jpeg",
            ".png",
            ".webp",
            ".gif",

        }:

            st.image(

                uploaded_file,

                caption=uploaded_file.name,

                use_container_width=True,
            )

            st.success(
                "🖼️ Image ready. "
                "Ask a question about the uploaded image."
            )

        # ----------------------------------------------------
        # VIDEO
        # ----------------------------------------------------

        if extension in {

            ".mp4",
            ".mov",
            ".m4v",
            ".avi",
            ".mkv",

        }:

            try:

                st.video(
                    uploaded_file
                )

            except Exception:
                pass

            if source:

                video_segments = (
                    source.get(
                        "video_segments",
                        []
                    )
                )

                transcript = (
                    source.get(
                        "uploaded_context",
                        ""
                    )
                )

                st.success(
                    "🎬 Video transcription completed."
                )

                st.caption(

                    f"Detected language: "
                    f"{LANGUAGE_NAMES.get(source.get('video_language', 'en'), 'English')}"
                )

                with st.expander(
                    "🎬 View video transcript with timestamps"
                ):

                    st.text_area(

                        "Transcript",

                        transcript,

                        height=300,

                        disabled=True,
                    )

        # ----------------------------------------------------
        # EXTRACTED TEXT
        # ----------------------------------------------------

        if source:

            extracted_context = (
                source.get(
                    "uploaded_context",
                    ""
                )
            )

            if extracted_context:

                with st.expander(
                    "📄 View extracted source text"
                ):

                    preview = extracted_context

                    if len(preview) > 6000:

                        preview = (
                            preview[:6000]
                            +
                            "\n\n...[truncated]"
                        )

                    st.text_area(

                        "Extracted content",

                        preview,

                        height=250,

                        disabled=True,
                    )

    # ========================================================
    # CURRENT SOURCE STATUS
    # ========================================================

    source = (
        st.session_state.source_data
    )

    if source:

        source_type = source.get(
            "source_type"
        )

        if source_type == "video":

            st.info(

                "🎬 Video source ready. "
                "Ask a question about any part of the video. "
                "The relevant frame and timestamp will be shown."
            )

        elif source_type == "pdf":

            st.info(

                "📄 PDF source ready. "
                "Ask about a section and the relevant page "
                "screenshot will be shown with the answer."
            )

        elif source_type in {
            "docx",
            "pptx"
        }:

            st.info(

                "📄 Office document ready. "
                "Relevant page/slide screenshot will be shown "
                "when available."
            )

        elif source_type == "website":

            st.info(

                "🌐 Website ready. "
                "Ask about a section and the relevant "
                "website screenshot will be shown."
            )

        elif source_type == "image":

            st.info(

                "🖼️ Image source ready. "
                "Ask anything about the uploaded image."
            )

        elif source_type == "text":

            st.info(

                "📄 Text source ready. "
                "Ask a question about the uploaded content."
            )

    # ========================================================
    # CHAT HISTORY
    # ========================================================

    render_messages()

    # ========================================================
    # VOICE INPUT
    # ========================================================

    st.subheader(
        "🎤 Voice Input"
    )

    st.caption(
        "Speak → Edit → Send to AI → "
        "Text + Voice Answer"
    )

    recorder_key = (
        f"recorder_"
        f"{st.session_state.recorder_key}"
    )

    audio_value = st.audio_input(

        "Click microphone and speak",

        sample_rate=16000,

        key=recorder_key,
    )

    # ========================================================
    # PROCESS NEW RECORDING
    # ========================================================

    if audio_value is not None:

        if not st.session_state.voice_ready:

            st.audio(
                audio_value,
                format="audio/wav"
            )

            if st.button(

                "🎧 Transcribe Voice",

                use_container_width=True,

                type="primary",
            ):

                input_path = None

                converted_path = None

                try:

                    input_path = os.path.join(

                        tempfile.gettempdir(),

                        f"recording_"
                        f"{uuid.uuid4().hex}.wav"
                    )

                    with open(

                        input_path,

                        "wb"

                    ) as audio_file:

                        audio_file.write(
                            audio_value.getvalue()
                        )

                    with st.spinner(
                        "🔄 Preparing audio..."
                    ):

                        converted_path = (
                            convert_audio_to_wav(
                                input_path
                            )
                        )

                    with st.spinner(

                        f"⏳ Loading Whisper "
                        f"{WHISPER_MODEL_NAME}..."

                    ):

                        whisper_model = (
                            load_whisper()
                        )

                    with st.spinner(
                        "🎧 Understanding your voice..."
                    ):

                        (
                            language_code,
                            transcript,
                        ) = transcribe_audio(

                            whisper_model,

                            converted_path,
                        )

                    if not transcript:

                        st.warning(

                            "⚠️ No speech detected. "
                            "Please record again."
                        )

                    else:

                        st.session_state.pending_transcript = (
                            transcript
                        )

                        st.session_state.pending_language = (
                            language_code
                        )

                        st.session_state.voice_ready = (
                            True
                        )

                        st.session_state.recorder_key += 1

                        st.rerun()

                except Exception as e:

                    st.error(
                        f"❌ Voice processing failed: {e}"
                    )

                    print(
                        "VOICE PROCESSING ERROR:",
                        repr(e)
                    )

                finally:

                    for path in [

                        input_path,
                        converted_path,

                    ]:

                        if path:

                            try:

                                if os.path.exists(
                                    path
                                ):

                                    os.remove(
                                        path
                                    )

                            except Exception:
                                pass

    # ========================================================
    # EDITABLE VOICE QUESTION
    # ========================================================

    if st.session_state.voice_ready:

        st.divider()

        st.subheader(
            "✏️ Edit Your Voice Question"
        )

        detected_language = (
            st.session_state.pending_language
        )

        detected_name = (
            LANGUAGE_NAMES.get(
                detected_language,
                detected_language
            )
        )

        st.info(
            f"🌍 Detected language: "
            f"{detected_name}"
        )

        edited_question = st.text_area(

            "Edit the transcription before sending to AI",

            value=(
                st.session_state.pending_transcript
            ),

            height=160,

            key="voice_question_editor",

            help=(
                "Correct transcription mistakes "
                "before sending."
            ),
        )

        col1, col2 = st.columns(2)

        with col1:

            send_voice_question = st.button(

                "🚀 Send Edited Question to AI",

                use_container_width=True,

                type="primary",
            )

        with col2:

            cancel_voice_question = st.button(

                "❌ Cancel",

                use_container_width=True,
            )

        if cancel_voice_question:

            st.session_state.pending_transcript = ""

            st.session_state.pending_language = "en"

            st.session_state.voice_ready = False

            if "voice_question_editor" in (
                st.session_state
            ):

                del st.session_state[
                    "voice_question_editor"
                ]

            st.session_state.recorder_key += 1

            st.rerun()

        if send_voice_question:

            final_question = (
                edited_question.strip()
            )

            if not final_question:

                st.warning(
                    "⚠️ Please enter a question."
                )

            else:

                final_language = (
                    detect_text_language(
                        final_question
                    )
                )

                original_language = (
                    st.session_state.pending_language
                )

                if (

                    final_language == "en"

                    and

                    original_language != "en"

                    and

                    detect_script_language(
                        final_question
                    ) is None

                ):

                    final_language = (
                        original_language
                    )

                if final_language not in LANGUAGE_NAMES:

                    final_language = "en"

                process_user_message(

                    user_message=
                    final_question,

                    language_code=
                    final_language,
                )

                st.session_state.pending_transcript = ""

                st.session_state.pending_language = "en"

                st.session_state.voice_ready = False

                if "voice_question_editor" in (
                    st.session_state
                ):

                    del st.session_state[
                        "voice_question_editor"
                    ]

                st.session_state.recorder_key += 1

                st.rerun()

    # ========================================================
    # PROCESS QUESTION THAT WAS WAITING FOR LOGIN
    # ========================================================

    pending_website_question = st.session_state.get("pending_website_question", "").strip()
    if pending_website_question and st.session_state.get("website_authenticated", False):
        st.session_state.pending_website_question = ""
        process_user_message(
            user_message=pending_website_question,
            language_code=detect_text_language(pending_website_question),
        )
        st.rerun()

    # ========================================================
    # TEXT CHAT
    # ========================================================

    user_text = st.chat_input(
        "Type your message here..."
    )

    if user_text:

        language_code = detect_text_language(
            user_text
        )

        print("=" * 70)
        print("TYPED MESSAGE")
        print("=" * 70)

        print(
            f"Text     : {user_text}"
        )

        print(
            f"Language : {language_code}"
        )

        print("=" * 70)

        process_user_message(

            user_message=user_text,

            language_code=language_code,
        )

        st.rerun()


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":

    main()